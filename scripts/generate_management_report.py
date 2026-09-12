#!/usr/bin/env python
"""Generate management briefing Word document for City Point ERP/Portal."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "docs" / "City_Point_ERP_Rehberlik_Hesabati.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_custom(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=RGBColor(0x12, 0x3B, 0x68))
    return h


def add_para(doc, text, *, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(0x12, 0x3B, 0x68))
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("CITY POINT BAKU")
    set_run_font(r, size=22, bold=True, color=RGBColor(0x12, 0x3B, 0x68))

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("Biznes Mərkəzi İdarəetmə Platforması")
    set_run_font(r2, size=16, bold=True, color=RGBColor(0x0C, 0x27, 0x45))

    t3 = doc.add_paragraph()
    t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = t3.add_run("Resident Portal + Staff ERP")
    set_run_font(r3, size=12, color=RGBColor(0xC8, 0xA7, 0x6B))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(18)
    rm = meta.add_run(
        f"Rəhbərlik üçün funksional hesabat\n"
        f"Görülən işlər · Mövcud imkanlar · Gələcək inkişaf\n"
        f"Tarix: {date.today().strftime('%d.%m.%Y')}"
    )
    set_run_font(rm, size=11)

    add_para(
        doc,
        "Bu sənəd City Point Baku üçün hazırlanan rəqəmsal idarəetmə platformasının "
        "cari vəziyyətini, əhatə olunan funksionallığı, tamamlanmış işləri və "
        "növbəti mərhələ planını rəhbərlik üçün aydın və ətraflı şəkildə təqdim edir.",
        space_after=14,
    )

    # 1
    add_heading_custom(doc, "1. Layihənin mahiyyəti və məqsədi", 1)
    add_para(
        doc,
        "City Point platforması biznes mərkəzinin gündəlik əməliyyatlarını bir mərkəzdə "
        "toplamaq üçün yaradılmışdır. Məqsəd icarəçi (rezident) şirkətlərlə bina "
        "idarəetmə komandası arasında sürətli, izlənilən və şəffaf əlaqə yaratmaq, "
        "xidmət müraciətlərini (ticket) standartlaşdırmaq, reception proseslərini "
        "rəqəmsallaşdırmaq və rəhbərlik üçün operativ vəziyyəti görünən etməkdir.",
    )
    add_para(doc, "Platforma iki əsas istifadəçi səthindən ibarətdir:", bold=True)
    add_bullet(
        doc,
        "Resident Portal — icarəçi şirkətlərin ofis əməkdaşları üçün: müraciət açmaq, "
        "status izləmək, sahə/elan/sənəd/qonaq məlumatlarına baxmaq.",
    )
    add_bullet(
        doc,
        "Staff ERP — City Point əməkdaşları (reception, service desk, property/FM, "
        "rəhbərlik, admin) üçün: operativ panel, qonaq qeydiyyatı, ticket idarəetməsi, "
        "əmlak və rezident baxışı, hesabatlar.",
    )
    add_para(
        doc,
        "Nəticə etibarilə platforma “kağız/telefon/WhatsApp üzərindən dağınıq "
        "kommunikasiya” modelindən “vahid rəqəmsal iş axını + status + SLA” modelinə "
        "keçidi təmin edir.",
    )

    # 2
    add_heading_custom(doc, "2. Görülən işlər (tamamlanmış mərhələ)", 1)
    add_para(
        doc,
        "Aşağıdakılar hazırkı mərhələdə reallaşdırılmış əsas nəticələrdir. "
        "Bunlar demo mühitində işlək vəziyyətdədir və rəhbərlik/əməkdaşlar tərəfindən "
        "yoxlanıla bilər.",
    )

    add_heading_custom(doc, "2.1. Arxitektura və infrastruktur", 2)
    add_bullet(doc, "Django 5 əsaslı monolit veb tətbiq (server-side templates).")
    add_bullet(doc, "PostgreSQL məlumat bazası.")
    add_bullet(doc, "Docker Compose ilə inkişaf (dev) və istehsal (prod) hazırlığı.")
    add_bullet(doc, "Sessiya əsaslı giriş (email + şifrə), rol əsaslı yönləndirmə və səhifə icazələri.")
    add_bullet(doc, "Demo məlumatların avtomatik yüklənməsi (seed) — real ssenarilərə yaxın nümunələr.")

    add_heading_custom(doc, "2.2. İstifadəçi interfeysi və brend", 2)
    add_bullet(doc, "Vahid City Point dizayn sistemi (navy / soft gold tokenlər, sidebar, kartlar, badge-lər).")
    add_bullet(doc, "Portal və ERP üçün eyni vizual dil — öyrənmə əyrisini azaldır.")
    add_bullet(doc, "Login səhifəsi, naviqasiya, KPI kartları, status və prioritet nişanları.")
    add_bullet(doc, "Responsive iş masası təcrübəsi (desktop prioritetli, mobilə uyğun əsas layout).")

    add_heading_custom(doc, "2.3. Çoxdillilik (i18n)", 2)
    add_bullet(doc, "Dəstəklənən dillər: Azərbaycan (əsas), İngilis, Rus.")
    add_bullet(doc, "Header və login səhifəsində dil seçici — eyni səhifədə qalmaqla dil dəyişməsi.")
    add_bullet(doc, "Menyu, düymələr, statuslar və əsas UI mətnləri tərcümə kataloqlarına salınıb.")

    add_heading_custom(doc, "2.4. Biznes modullarının qurulması", 2)
    add_bullet(doc, "Rol modeli: Admin, Rəhbərlik, Reception, Service Desk, Property/FM, Rezident.")
    add_bullet(doc, "Resident Portal: əsas səhifə, sahələr, müraciətlər, bildirişlər, əməkdaş girişi, qonaqlar, elanlar, sənədlər.")
    add_bullet(doc, "Staff ERP: idarə paneli, reception, ticketlər, sahələr, rezidentlər, sənədlər, hesabatlar.")
    add_bullet(doc, "Ticket iş axını + SLA siyasəti + status tarixçəsi + daxili söhbət.")
    add_bullet(doc, "Reception: qonaq qeydiyyatı, check-in / check-out.")
    add_bullet(doc, "Əmlak baxışı: bina / mərtəbə / sahə strukturu və real saylar (süni statistikalar olmadan).")

    add_heading_custom(doc, "2.5. Təhlükəsizlik və əməliyyat intizamı", 2)
    add_bullet(doc, "Rolə görə menyu və səhifə girişi məhdudlaşdırılıb.")
    add_bullet(doc, "Rezident yalnız öz şirkətinin məlumatlarını görür.")
    add_bullet(doc, "CSRF qoruması, sessiya cookie, Django Admin ilə master-data idarəetmə imkanı.")

    # 3
    add_heading_custom(doc, "3. Mövcud funksionallıq — ətraflı izah", 1)

    add_heading_custom(doc, "3.1. Giriş və rollər", 2)
    add_para(
        doc,
        "İstifadəçi email və şifrə ilə daxil olur. Sistem avtomatik olaraq rolə görə "
        "Portal və ya ERP-yə yönləndirir. Bu, yanlış portalda işləmə riskini aradan qaldırır "
        "və hər komandaya yalnız ona lazım olan alətləri göstərir.",
    )
    add_table(
        doc,
        ["Rol", "Əsas təyinat", "Əsas giriş səthi"],
        [
            ["Admin", "Tam idarəetmə, konfiqurasiya", "ERP + Django Admin"],
            ["Rəhbərlik", "Ümumi nəzarət və hesabat", "ERP (bütün əsas bölmələr)"],
            ["Reception", "Qonaq axını və qarşılama", "ERP → Reception"],
            ["Service Desk", "Xidmət müraciətləri", "ERP → Ticketlər"],
            ["Property / FM", "Sahə və texniki baxış", "ERP → Sahələr + Ticketlər"],
            ["Rezident istifadəçi", "Şirkət ofisi əməkdaşı", "Resident Portal"],
        ],
    )

    add_heading_custom(doc, "3.2. Resident Portal (icarəçi tərəfi)", 2)
    add_para(
        doc,
        "Portal rezident şirkətin “rəqəmsal pəncərəsidir”. Məqsəd: binaya müraciət etmək "
        "üçün telefon zəngi və ya şəxsi əlaqəyə ehtiyacı azaltmaq, eyni zamanda prosesi "
        "şəffaf etmək.",
    )
    add_bullet(doc, "Əsas səhifə — açıq müraciətlər, elanlar və ümumi xülasə.")
    add_bullet(doc, "Sahələrim — şirkətə bağlı ofis/sahə məlumatları.")
    add_bullet(
        doc,
        "Yeni müraciət — kateqoriya, sahə, prioritet, təsvir və foto/sənəd əlavəsi ilə ticket yaratmaq.",
    )
    add_bullet(doc, "Aktiv müraciətlər / Tarixçə — açıq və bağlanmış müraciətlərin izlənməsi.")
    add_bullet(
        doc,
        "Müraciət detalı — status, SLA qalan vaxt, status tarixçəsi və service desk ilə mesajlaşma.",
    )
    add_bullet(doc, "Bildirişlər — ticket yeniləmələri və mesajlar barədə xəbərdarlıqlar.")
    add_bullet(doc, "Əməkdaşlar — gün ərzində giriş/çıxış (access) hadisələrinə baxış.")
    add_bullet(doc, "Qonaqlar — həmin günə qeydiyyatdan keçmiş qonaqlara baxış.")
    add_bullet(doc, "Elanlar və Sənədlər — bina/şirkət kommunikasiyası və sənəd kataloqu.")

    add_heading_custom(doc, "3.3. Staff ERP — İdarə paneli", 2)
    add_para(
        doc,
        "ERP dashboard əməliyyat komandası üçün “bir baxışda vəziyyət” verir: "
        "bugünkü qonaq sayı, açıq ticketlər, SLA riskində olanlar, gözləyən qonaqlar "
        "və son fəaliyyətlər. Bu, rəhbərlik və növbə rəhbərləri üçün operativ qərar "
        "verməni asanlaşdırır.",
    )

    add_heading_custom(doc, "3.4. Reception (qonaq idarəetməsi)", 2)
    add_para(
        doc,
        "Reception modulu qonaq axınını rəqəmsallaşdırır. Əməkdaş yeni qonağı qeyd edir, "
        "statusunu check-in / check-out ilə yeniləyir. Bu, təhlükəsizlik və müştəri "
        "xidməti üçün izlənilən jurnal yaradır və portalda rezident tərəfinə də "
        "görünürlük verir (baxış).",
    )
    add_bullet(doc, "Bugünkü qonaq siyahısı.")
    add_bullet(doc, "Yeni qonaq qeydiyyatı.")
    add_bullet(doc, "Check-in və check-out əməliyyatları.")

    add_heading_custom(doc, "3.5. Service Desk və Ticket / SLA", 2)
    add_para(
        doc,
        "Ticket sistemi platformanın ürəyidir. Hər müraciət unikal kod alır, prioritetə "
        "görə SLA müddəti təyin olunur, statuslar addım-addım irəliləyir və hər keçid "
        "tarixçədə saxlanılır. Bu, “kim nə vaxt nə etdi” sualına cavab verir və "
        "xidmət keyfiyyətini ölçülə bilən edir.",
    )
    add_para(doc, "Status axını:", bold=True)
    add_bullet(doc, "Göndərildi → Qəbul edildi → İcra olunur → Həll edildi")
    add_para(doc, "Prioritet və nümunəvi SLA (demo siyasət):", bold=True)
    add_bullet(doc, "Yüksək — qısa müddət (kritik hallar)")
    add_bullet(doc, "Normal — standart xidmət müddəti")
    add_bullet(doc, "Aşağı — planlaşdırıla bilən işlər")
    add_para(doc, "Service Desk əməkdaşı edə bilər:", bold=True)
    add_bullet(doc, "Bütün / açıq / SLA risk / mənim ticketlərim filtrləri.")
    add_bullet(doc, "Cavabdeh təyin etmək.")
    add_bullet(doc, "Statusu növbəti mərhələyə keçirmək.")
    add_bullet(doc, "Rezidentlə daxili söhbət aparmaq (mesajlar bildirişə çevrilir).")
    add_para(
        doc,
        "Demo kateqoriyalar: HVAC, Access Card, Təmizlik, Elektrik — real bina "
        "xidmətlərinə uyğun başlanğıc təsnifat.",
    )

    add_heading_custom(doc, "3.6. Əmlak (Property) və Rezidentlər", 2)
    add_para(
        doc,
        "Əmlak modulu bina strukturunu (mərtəbə, sahə) və doluluğu real məlumat əsasında "
        "göstərir. Süni “14 mərtəbə / 96% doluluq” kimi marketinq rəqəmləri istifadə "
        "edilmir — yalnız sistemdə olan faktiki qeydlər əks olunur. Bu, rəhbərlik üçün "
        "etibarlı operativ mənzərə yaradır.",
    )
    add_bullet(doc, "Sahə siyahısı və mərtəbə üzrə baxış.")
    add_bullet(doc, "Sahə detalı: ümumi məlumat, bağlı ticketlər, sənədlər, əməkdaşlar, aktivlər.")
    add_bullet(doc, "Rezident şirkət siyahısı və detalı: sahələr, ticketlər, əməkdaşlar, bugünkü qonaqlar.")

    add_heading_custom(doc, "3.7. Sənədlər, elanlar və hesabatlar", 2)
    add_bullet(
        doc,
        "Sənədlər — şirkətə və ümumi bina sənədlərinə kataloq baxışı (metadata; fayl "
        "mövcuddursa açmaq mümkündür).",
    )
    add_bullet(
        doc,
        "Elanlar — portalda oxunur; yaradılma hazırda Admin panel üzərindən aparılır "
        "(ERP-də ayrıca redaktor növbəti mərhələyə saxlanılıb).",
    )
    add_bullet(
        doc,
        "Hesabatlar — rəhbərlik üçün ilkin KPI xülasəsi və aktiv rezident keçidləri "
        "(dərin analitika və export növbəti mərhələdədir).",
    )

    # 4
    add_heading_custom(doc, "4. Demo giriş məlumatları (yoxlama üçün)", 1)
    add_para(
        doc,
        "İnkişaf mühitində sistem http://localhost:8000 ünvanında işləyir. "
        "Aşağıdakı hesablarla rollər üzrə ssenarilər nümayiş etdirilə bilər:",
    )
    add_table(
        doc,
        ["Rol", "Email", "Şifrə"],
        [
            ["Portal (ASBC)", "office@asbc.az", "asbc123"],
            ["Admin", "admin@citypoint.az", "admin123"],
            ["Reception", "reception@citypoint.az", "reception123"],
            ["Service Desk", "desk@citypoint.az", "desk123"],
            ["Property / FM", "fm@citypoint.az", "fm123"],
            ["Rəhbərlik", "manager@citypoint.az", "manager123"],
        ],
    )
    add_para(doc, "Əsas keçidlər:", bold=True)
    add_bullet(doc, "Giriş: http://localhost:8000/login/")
    add_bullet(doc, "Portal: http://localhost:8000/portal/")
    add_bullet(doc, "ERP: http://localhost:8000/erp/")

    # 5
    add_heading_custom(doc, "5. Gələcək inkişaf (roadmap)", 1)
    add_para(
        doc,
        "Hazırkı versiya operativ MVP-dir: əsas iş axınları işləyir, lakin biznes "
        "mərkəzinin tam rəqəmsal ekosisteminə çevrilmək üçün mərhələli genişlənmə "
        "tələb olunur. Aşağıdakı istiqamətlər prioritetləşdirilə bilər.",
    )

    add_heading_custom(doc, "5.1. Qısa müddət (0–3 ay) — operativ dərinləşmə", 2)
    add_bullet(doc, "ERP daxilində elan və sənəd yaratma/redaktə formları (Admin asılılığını azaltmaq).")
    add_bullet(doc, "Sənəd yükləmə workflow-u (versiya, tip, icazə).")
    add_bullet(doc, "Portalda qonaq əvvəlcədən qeydiyyat (pre-registration) sorğusu.")
    add_bullet(doc, "Hesabatlarda Excel/PDF export və sadə qrafiklər (SLA, ticket həcm, reception).")
    add_bullet(doc, "Email/SMS bildiriş şablonları (ticket status dəyişəndə).")
    add_bullet(doc, "Mobil uyğunluğun və istifadəçi təlimatının möhkəmləndirilməsi.")

    add_heading_custom(doc, "5.2. Orta müddət (3–9 ay) — əməliyyat genişlənməsi", 2)
    add_bullet(doc, "Texniki plan / DWG / Maintenance tablarının real məzmunla doldurulması.")
    add_bullet(doc, "Aktiv (asset) və podratçı (contractor) idarəetməsi — planlı təmir, iş sifarişi.")
    add_bullet(doc, "Anbar və satınalma (warehouse / purchasing) — ehtiyat hissə və təchizatçı axını.")
    add_bullet(doc, "CRM toxunuşları — rezident əlaqə tarixçəsi, razılaşma və xidmət keyfiyyəti qeydləri.")
    add_bullet(doc, "SLA siyasətlərinin kateqoriya üzrə çevik konfiqurasiyası və eskalasiya qaydaları.")
    add_bullet(doc, "Daha zəngin audit jurnalı və rol əsaslı hesabat paketləri.")

    add_heading_custom(doc, "5.3. Uzun müddət (9–18 ay) — inteqrasiya və maliyyə", 2)
    add_bullet(doc, "Access control (turniket/kart) sistemləri ilə canlı sinxronizasiya.")
    add_bullet(doc, "Mühasibat / billing / invoice — xidmət haqqı və icarə əlaqəli hesablaşma.")
    add_bullet(doc, "Xarici API (mobil tətbiq, partnyor sistemlər, BI alətləri).")
    add_bullet(doc, "İş axını avtomatlaşması (növbə, avtomatik təyinat, SLA breach alert).")
    add_bullet(doc, "İstehsal mühitində monitoring, ehtiyat nüsxə və yüksək əlçatanlıq.")

    add_heading_custom(doc, "5.4. Strateji baxış", 2)
    add_para(
        doc,
        "Uzunmüddətli vizyon: City Point platformasının biznes mərkəzinin “əməliyyat "
        "əməliyyat sistemi”nə (operating system) çevrilməsi — reception-dan service "
        "desk-ə, əmlakdan maliyyəyə qədər vahid məlumat və proses zənciri. Bu, "
        "xidmət keyfiyyətini artırır, əməkdaş vaxtını azaldır, rezident məmnuniyyətini "
        "ölçülə bilən edir və rəhbərlik üçün idarəetmə qərarlarını data əsaslı edir.",
    )

    # 6
    add_heading_custom(doc, "6. Biznes faydası (rəhbərlik üçün)", 1)
    add_bullet(doc, "Şəffaflıq — hər müraciətin statusu və tarixçəsi görünür.")
    add_bullet(doc, "SLA intizamı — gecikmələr risk və breach kimi erkən görünür.")
    add_bullet(doc, "Rolə uyğun iş yeri — hər komanda yalnız öz funksiyasını görür.")
    add_bullet(doc, "Rezident self-service — yükün bir hissəsi portalə keçir.")
    add_bullet(doc, "Çoxdillilik — AZ/EN/RU ilə beynəlxalq icarəçilərə uyğunluq.")
    add_bullet(doc, "Genişlənməyə açıq baza — anbar, CRM, maintenance, billing üçün təməl hazırdır.")

    # 7
    add_heading_custom(doc, "7. Məhdudiyyətlər və cari sərhədlər", 1)
    add_para(
        doc,
        "Dürüst hesabat üçün hazırkı versiyanın hələ əhatə etmədiyi məqamlar da "
        "qeyd olunmalıdır:",
    )
    add_bullet(doc, "Master-data (şirkət, sahə, elan və s. yaratma) əsasən Admin paneldən aparılır.")
    add_bullet(doc, "Sənəd kataloqunda bir çox qeyd metadata səviyyəsindədir; tam fayl arxivi hələ tamamlanmayıb.")
    add_bullet(doc, "Hesabatlar ilkin KPI səviyyəsindədir — dərin BI/export yoxdur.")
    add_bullet(doc, "Canlı turniket/ACS inteqrasiyası, email gateway, REST API hələ yoxdur.")
    add_bullet(doc, "Texniki plan / DWG / Maintenance bölmələri interfeysdə placeholder kimi saxlanılıb.")
    add_para(
        doc,
        "Bu məhdudiyyətlər layihənin uğursuzluğu deyil — MVP sərhədləridir. Növbəti "
        "sprintlərdə məhz bu boşluqlar doldurulmalıdır.",
    )

    # 8
    add_heading_custom(doc, "8. Tövsiyələr", 1)
    add_bullet(
        doc,
        "Rəhbərlik və əsas rollərlə (reception, desk, FM, bir rezident) canlı demo "
        "sessiyası keçirilsin — real rəy toplamaq üçün.",
    )
    add_bullet(
        doc,
        "Qısa müddətli prioritetlər təsdiqlənsin: sənəd/elan ERP formaları, portal "
        "qonaq pre-registration, hesabat export, email bildiriş.",
    )
    add_bullet(
        doc,
        "İstehsal (prod) mühiti üçün hosting, ehtiyat nüsxə və istifadəçi təlim "
        "planı müəyyənləşdirilsin.",
    )
    add_bullet(
        doc,
        "SLA müddətləri və kateqoriyalar City Point əməliyyat qaydalarına uyğun "
        "yenidən konfiqurasiya edilsin.",
    )

    # 9
    add_heading_custom(doc, "9. Nəticə", 1)
    add_para(
        doc,
        "City Point Resident Portal + Staff ERP platforması hazırkı mərhələdə "
        "işlək operativ MVP kimi təqdim olunur. Əsas dəyər zənciri — müraciət "
        "yaratma, status izləmə, SLA nəzarəti, reception qeydiyyatı, əmlak/rezident "
        "baxışı və çoxdilli interfeys — qurulub və nümayişə hazırdır.",
    )
    add_para(
        doc,
        "Gələcək mərhələlərdə sənəd və elan idarəetməsinin ERP-yə köçürülməsi, "
        "analitikanın gücləndirilməsi, maintenance/anbar/CRM genişlənməsi və "
        "xarici sistem inteqrasiyaları platformanı tam biznes mərkəzi idarəetmə "
        "həllinə çevirəcəkdir.",
    )
    add_para(
        doc,
        "Hazırkı baza uğurlu davam üçün yetərlidir: texniki təməl möhkəmdir, "
        "istifadəçi rolləri aydın bölünüb, UI vahiddir və biznes prosesləri "
        "ölçülə bilən formaya salınıb.",
        bold=True,
    )

    footer = doc.add_paragraph()
    footer.paragraph_format.space_before = Pt(24)
    fr = footer.add_run(
        "— Sənəd City Point Baku rəqəmsal platforması üzrə daxili rəhbərlik brifinqi üçün hazırlanmışdır. —"
    )
    set_run_font(fr, size=9, color=RGBColor(0x66, 0x66, 0x66))
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
