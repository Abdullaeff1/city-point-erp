#!/usr/bin/env python3
"""Generate AlmaLinux 10 City Point ERP deployment guide as Word (.docx)."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "deployment" / "CityPoint_AlmaLinux10_Deploy_Rehberi.docx"


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color=RGBColor(0x1F, 0x49, 0x7D))
    return h


def add_para(doc, text, *, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    if level:
        p.paragraph_format.left_indent = Cm(0.75 * level)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        set_run_font(run, bold=True, size=10)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            run = cells[c_i].paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("CITY POINT ERP / PORTAL")
    set_run_font(r, size=22, bold=True, color=RGBColor(0x1F, 0x49, 0x7D))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("AlmaLinux 10 üzərində quraşdırma və istismar rəhbəri")
    set_run_font(r, size=14, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        f"Versiya: 1.0  ·  Tarix: {date.today().isoformat()}  ·  "
        "Stack: Django 5 + PostgreSQL 16 + Docker Compose  ·  Model: LAN / air-gapped"
    )
    set_run_font(r, size=9, color=RGBColor(0x64, 0x74, 0x8B))

    add_para(
        doc,
        "Bu sənəd City PointBaku MMC daxili ERP və Resident Portal sisteminin "
        "AlmaLinux 10 serverində işə salınması üçün lazımi resursları, NetAdmin / "
        "SysAdmin / ERP ops işlərini, şəbəkə, Docker, AxTrax inteqrasiyası, backup "
        "və qəbul checklist-lərini əhatə edir.",
        size=11,
    )

    # ---- 1 ----
    add_heading(doc, "1. Məqsəd və arxitektura", 1)
    add_para(
        doc,
        "Sistem tək (və ya klaster) AlmaLinux 10 hostunda Docker Compose ilə işləyir. "
        "Brauzerlər bina LAN-ından ERP/Portal-a qoşulur. AxTraxNG MS SQL oxuma (RO) "
        "LAN üzərindən həyata keçirilir. İstismar zamanı İnternet tələb olunmur.",
    )
    add_para(doc, "Komponentlər:", bold=True)
    add_bullet(doc, "web — Django 5 (Gunicorn), Python 3.12, AxTrax poller prosesi")
    add_bullet(doc, "db — PostgreSQL 16 (Docker volume)")
    add_bullet(doc, "media / staticfiles — Docker volume-lar")
    add_bullet(doc, "opsional: Nginx reverse proxy + daxili TLS")
    add_para(doc, "Şəbəkə modeli (mətn):", bold=True)
    add_code(
        doc,
        "[ Resident / Staff PC ] → LAN switch → AlmaLinux Docker host (:8000 və ya :80/:443)\n"
        "                                         ↓\n"
        "                              AxTrax MS SQL (LAN, RO, adətən :1433)",
    )
    add_para(doc, "Repo referansları:", bold=True)
    add_bullet(doc, "bin/prod/docker-compose.yml — prod compose")
    add_bullet(doc, "bin/prod/.env.example — mühit dəyişənləri")
    add_bullet(doc, "docs/deployment/air-gapped-lan.md — LAN/air-gap qaydaları")
    add_bullet(doc, "docs/security/postgres-roles.md — DB rolları")
    add_bullet(doc, "docs/backup/phase1-backup-restore.md — backup/restore")

    # ---- 2 ----
    add_heading(doc, "2. Server resursları", 1)
    add_heading(doc, "2.1. Hardware / VM", 2)
    add_table(
        doc,
        ["Resurs", "Minimum (pilot)", "Tövsiyə (istismar)", "Qeyd"],
        [
            ["CPU", "4 vCPU", "8 vCPU", "Poller + Gunicorn + Postgres"],
            ["RAM", "8 GB", "16 GB", "DB cache və eyni hostda media"],
            ["Disk", "100 GB SSD", "200+ GB SSD", "DB + media + lokal backup"],
            ["OS", "AlmaLinux 10 x86_64", "AlmaLinux 10 x86_64", "RHEL 10 uyğun"],
            ["Şəbəkə", "1 Gbit LAN", "1 Gbit + statik IP", "AxTrax-a marşrut"],
            ["Backup media", "NAS / ikinci disk", "NAS + offsite surət", "pg_dump + media"],
        ],
    )
    add_heading(doc, "2.2. Proqram resursları (siyahı)", 2)
    add_bullet(doc, "Docker Engine + Docker Compose plugin (və ya Podman uyğun stack — layihə Docker Compose sənədləşdirilib)")
    add_bullet(doc, "firewalld (və ya nftables)")
    add_bullet(doc, "chrony (daxili NTP)")
    add_bullet(doc, "Opsional: nginx, cert (daxili CA)")
    add_bullet(doc, "Repo arxivi (git zip) + docker image tar-ları (air-gap üçün)")
    add_bullet(doc, "AxTrax ReadOnly SQL hesabı və şəbəkə girişi")

    # ---- 3 ----
    add_heading(doc, "3. Əvvəlcədən toplanmalı məlumatlar", 1)
    add_para(doc, "NetAdmin / IT bu siyahını doldurmadan quraşdırma tamamlanmamalıdır:", bold=True)
    add_table(
        doc,
        ["#", "Məlumat", "Nümunə / qeyd"],
        [
            ["1", "Server hostname", "erp.citypoint.local"],
            ["2", "Server statik IP", "10.x.x.x"],
            ["3", "Portal/ERP URL", "http://erp.citypoint.local:8000"],
            ["4", "AxTrax SQL host/port/DB", "172.31.104.10:1433 / AxTrax1"],
            ["5", "AxTrax RO user/password", "ReadOnlyerp / (secret store)"],
            ["6", "Mail yolu", "Daxili SMTP və ya CP_OFFLINE=1 + link kopyalama"],
            ["7", "DJANGO_SECRET_KEY", "Güclü təsadüfi (chatə yazılmasın)"],
            ["8", "Postgres şifrələri", "cp_migrator / cp_app / cp_readonly"],
            ["9", "Backup yolu", "/backup/citypoint və ya NAS share"],
            ["10", "NTP mənbəyi", "Domain DC / daxili chrony"],
            ["11", "Firewall qaydaları", "8000 (və ya 80/443) + AxTrax 1433 çıxış"],
            ["12", "Pilot şirkətlər", "1–2 rezident + portal kontaktlar"],
        ],
    )

    # ---- 4 ----
    add_heading(doc, "4. AlmaLinux 10 OS hazırlığı — görüləcək işlər", 1)
    add_heading(doc, "4.1. İlkin quraşdırma", 2)
    add_bullet(doc, "AlmaLinux 10 Minimal (və ya Server) quraşdırın (x86_64).")
    add_bullet(doc, "Vaxt zonasını Asia/Baku (və ya City Point siyasəti) təyin edin.")
    add_bullet(doc, "Root/sudo admin hesabı yaradın; ops üçün citypoint istifadəçisi tövsiyə olunur.")
    add_bullet(doc, "Sistem yeniləmələrini texniki pəncərədə edin (air-gap-də: daxili mirror / USB RPM).")

    add_heading(doc, "4.2. İstifadəçi və qruplar", 2)
    add_code(
        doc,
        "sudo useradd -m -s /bin/bash citypoint\n"
        "sudo usermod -aG wheel citypoint\n"
        "# Docker quraşdırıldıqdan sonra:\n"
        "sudo usermod -aG docker citypoint",
    )

    add_heading(doc, "4.3. Vaxt sinxronizasiyası (mütləq)", 2)
    add_para(
        doc,
        "AxTrax hadisə vaxtları və Excel hesabatları düzgün saat tələb edir. "
        "Daxili NTP (domain controller və ya chrony) istifadə edin — public NTP air-gap-də olmaya bilər.",
    )
    add_code(
        doc,
        "sudo dnf install -y chrony\n"
        "sudo systemctl enable --now chronyd\n"
        "timedatectl status",
    )

    add_heading(doc, "4.4. Firewall", 2)
    add_code(
        doc,
        "sudo dnf install -y firewalld\n"
        "sudo systemctl enable --now firewalld\n"
        "# ERP/Portal (birbaşa Docker publish)\n"
        "sudo firewall-cmd --permanent --add-port=8000/tcp\n"
        "# Əgər Nginx: 80/443\n"
        "# sudo firewall-cmd --permanent --add-service=http\n"
        "# sudo firewall-cmd --permanent --add-service=https\n"
        "sudo firewall-cmd --reload",
    )
    add_para(
        doc,
        "Çıxış: hostdan AxTrax MS SQL (TCP 1433) və daxili SMTP (əgər varsa) icazəli olmalıdır. "
        "İnternet uplink opsionaldır; air-gap go-live üçün runtime İnternet lazım deyil.",
    )

    add_heading(doc, "4.5. SELinux", 2)
    add_para(
        doc,
        "AlmaLinux-da SELinux adətən Enforcing-dir. Docker volume və bind-mount "
        "problemlərində konteyner loglarına baxın. Prod-da SELinux-u söndürmək tövsiyə olunmur; "
        "lazım gələrsə volume üçün uyğun kontekst və ya :z/:Z mount seçimləri sənədləşdirilsin.",
    )
    add_code(doc, "getenforce\nsudo ausearch -m AVC -ts recent | tail")

    add_heading(doc, "4.6. Docker Engine + Compose", 2)
    add_para(
        doc,
        "Layihə rəsmi olaraq Docker Compose ilə sənədləşdirilib. AlmaLinux 10-da Docker CE "
        "və ya daxili siyasətə uyğun konteyner runtime quraşdırın. İnternet yoxdursa RPM-ləri "
        "offline köçürün.",
    )
    add_code(
        doc,
        "# İnternet olan mühitdə (nümunə axını — versiyalar IT mirror-a görə dəyişə bilər):\n"
        "sudo dnf -y install dnf-plugins-core\n"
        "# Docker rəsmi repo əlavə edin və ya daxili mirror istifadə edin\n"
        "sudo dnf -y install docker-ce docker-ce-cli containerd.io docker-compose-plugin\n"
        "sudo systemctl enable --now docker\n"
        "docker version\n"
        "docker compose version",
    )

    # ---- 5 ----
    add_heading(doc, "5. Air-gap / offline paket hazırlığı", 1)
    add_para(
        doc,
        "Go-live zamanı serverdə docker pull / pip install İNTERNETƏ bağlı olmamalıdır. "
        "Image və kod əvvəlcədən hazırlanır.",
    )
    add_heading(doc, "5.1. İnternet olan build maşınında", 2)
    add_code(
        doc,
        "# Repoyu klonlayın / zip açın\n"
        "cd city-point-system\n"
        "docker pull postgres:16-alpine\n"
        "docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env build\n"
        "\n"
        "# Image-ləri saxlayın\n"
        "docker save postgres:16-alpine -o postgres_16_alpine.tar\n"
        "docker images   # citypoint prod web image adını qeyd edin\n"
        "docker save <web_image:tag> -o citypoint_web.tar",
    )
    add_bullet(doc, "USB/NAS-a: postgres_16_alpine.tar, citypoint_web.tar, repo zip, bin/prod/.env.example")
    add_bullet(doc, "Secret-ləri (.env real) ayrı daşıyın; chatə/repo-ya yazmayın")

    add_heading(doc, "5.2. AlmaLinux serverində", 2)
    add_code(
        doc,
        "sudo mkdir -p /opt/citypoint /opt/citypoint-images\n"
        "# tar və repo-nu köçürün\n"
        "docker load -i /opt/citypoint-images/postgres_16_alpine.tar\n"
        "docker load -i /opt/citypoint-images/citypoint_web.tar\n"
        "cd /opt/citypoint",
    )

    # ---- 6 ----
    add_heading(doc, "6. Tətbiqin quraşdırılması (addım-addım)", 1)
    add_heading(doc, "6.1. Kataloq və .env", 2)
    add_code(
        doc,
        "cd /opt/citypoint/bin/prod\n"
        "cp .env.example .env\n"
        "chmod 600 .env\n"
        "nano .env   # və ya vim",
    )
    add_para(doc, "Mütləq .env sahələri:", bold=True)
    add_table(
        doc,
        ["Dəyişən", "Tövsiyə dəyər"],
        [
            ["CP_ENV", "production"],
            ["CP_OFFLINE", "1 (SMTP yoxdursa) və ya 0 + daxili SMTP"],
            ["DJANGO_DEBUG", "0"],
            ["SEED_DEMO", "0"],
            ["DJANGO_SECRET_KEY", "güclü təsadüfi"],
            ["DJANGO_ALLOWED_HOSTS", "erp hostname, IP"],
            ["DJANGO_CSRF_TRUSTED_ORIGINS", "http://hostname (və ya https://)"],
            ["POSTGRES_*", "db adı, cp_app user, şifrə"],
            ["AXTRAX_POLL_ON_START", "1"],
            ["AXTRAX_POLL_INTERVAL", "15"],
            ["TURNSTILE_MSSQL_*", "AxTrax LAN parametrləri + password"],
        ],
    )

    add_heading(doc, "6.2. PostgreSQL rolları", 2)
    add_para(
        doc,
        "İlk qaldırmadan sonra superuser ilə cp_migrator, cp_app, cp_readonly yaradın "
        "(docs/security/postgres-roles.md). Migrate cp_migrator ilə, runtime POSTGRES_USER=cp_app.",
    )
    add_code(
        doc,
        "CREATE ROLE cp_migrator LOGIN PASSWORD '...';\n"
        "CREATE ROLE cp_app LOGIN PASSWORD '...';\n"
        "CREATE ROLE cp_readonly LOGIN PASSWORD '...';\n"
        "GRANT CONNECT ON DATABASE citypoint TO cp_migrator, cp_app, cp_readonly;",
    )

    add_heading(doc, "6.3. Compose işə salma", 2)
    add_code(
        doc,
        "cd /opt/citypoint\n"
        "docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env up -d\n"
        "# İlk dəfə build lazımdırsa (İnternet/build maşını):\n"
        "# docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env up -d --build\n"
        "\n"
        "docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env ps\n"
        "docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env logs -f web",
    )
    add_para(
        doc,
        "Prod compose web-i hostda 8000 portuna yayınlayır (bin/prod/docker-compose.yml). "
        "Dev mühitində 3001 istifadə olunur — qarışdırmayın.",
    )

    add_heading(doc, "6.4. Migrate və statik", 2)
    add_para(
        doc,
        "bin/prod/entrypoint.py adətən migrate/collectstatic edir. Əlavə əl ilə:",
    )
    add_code(
        doc,
        "docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env exec web \\\n"
        "  python manage.py migrate\n"
        "docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env exec web \\\n"
        "  python manage.py collectstatic --noinput",
    )

    add_heading(doc, "6.5. İlk istifadəçilər", 2)
    add_code(
        doc,
        "# Superuser / admin (siyasətə görə)\n"
        "docker compose ... exec web python manage.py createsuperuser\n"
        "\n"
        "# Rezident portal invite (public signup yoxdur)\n"
        "docker compose ... exec web python manage.py invite_portal_user \\\n"
        "  --email=office@tenant.az --company=<slug> --first-name=... --last-name=... \\\n"
        "  --base-url=http://erp.citypoint.local:8000 --send-email\n"
        "# CP_OFFLINE=1 olduqda linki admin kopyalayıb təhvil verir",
    )

    # ---- 7 ----
    add_heading(doc, "7. AxTraxNG inteqrasiyası (Linux)", 1)
    add_para(
        doc,
        "Windows Scheduled Task məcburi deyil: AXTRAX_POLL_ON_START=1 olduqda poller "
        "web konteyneri ilə birlikdə işləyir. Hostdan AxTrax MS SQL-ə TCP 1433 çatmalıdır.",
    )
    add_bullet(doc, "TURNSTILE_MSSQL_HOST / PORT / DATABASE / USER / PASSWORD .env-də")
    add_bullet(doc, "Status: python manage.py axtrax_sync_status")
    add_bullet(doc, "Tarixi doldurma: python manage.py backfill_axtrax_events --from=2026-09-01")
    add_bullet(doc, "People sync (daha az tez-tez): sync_axtrax_people axını — runbook-a baxın")
    add_code(
        doc,
        "docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env exec web \\\n"
        "  python manage.py axtrax_sync_status",
    )

    # ---- 8 ----
    add_heading(doc, "8. Backup və bərpa", 1)
    add_para(doc, "Backup etibarlı sayılmır, əgər restore drill edilməyibsə.", bold=True)
    add_heading(doc, "8.1. Nə ehtiyatlanır", 2)
    add_bullet(doc, "PostgreSQL citypoint dump")
    add_bullet(doc, "media volume (sənəd yükləmələri)")
    add_bullet(doc, "AxTrax MS SQL ERP backup-ına daxil deyil (ayrı sistem)")

    add_heading(doc, "8.2. Gündəlik dump nümunəsi (cron)", 2)
    add_code(
        doc,
        "# /etc/cron.d/citypoint-backup\n"
        "15 2 * * * root cd /opt/citypoint && \\\n"
        "  docker compose -f bin/prod/docker-compose.yml --env-file bin/prod/.env exec -T db \\\n"
        "  pg_dump -U \"$POSTGRES_USER\" \"$POSTGRES_DB\" | gzip > /backup/citypoint/db-$(date +\\%Y\\%m\\%d).sql.gz\n"
        "\n"
        "# Media (nümunə)\n"
        "30 2 * * * root tar czf /backup/citypoint/media-$(date +\\%Y\\%m\\%d).tar.gz -C /var/lib/docker/volumes ...",
    )
    add_para(doc, "Saxlama: 14 günlük + 8 həftəlik (City Point IT ilə razılaşdırın).")
    add_heading(doc, "8.3. Restore drill", 2)
    add_bullet(doc, "Dump-u ayrı DB adına bərpa edin (citypoint_restore_test)")
    add_bullet(doc, "Login + ticket siyahısı + AccessEvent sayını yoxlayın")
    add_bullet(doc, "Nəticəni ops jurnalına yazın")

    # ---- 9 ----
    add_heading(doc, "9. Reverse proxy (opsional, tövsiyə)", 1)
    add_para(
        doc,
        "Brauzerlərə birbaşa :8000 vermək əvəzinə Nginx ilə :80 (və daxili HTTPS) istifadə edilə bilər. "
        "DJANGO_CSRF_TRUSTED_ORIGINS və ALLOWED_HOSTS proxy hostname-ə uyğun yenilənməlidir. "
        "TLS üçün daxili CA sertifikatı istifadə olunur (Let's Encrypt air-gap-də yoxdur).",
    )

    # ---- 10 ----
    add_heading(doc, "10. Təhlükəsizlik və istismar qaydaları", 1)
    add_bullet(doc, "SEED_DEMO=0 — demo parollar prod-da olmamalıdır")
    add_bullet(doc, ".env chmod 600; yalnız ops qrupu oxusun")
    add_bullet(doc, "CDN yox — statik Whitenoise /staticfiles")
    add_bullet(doc, "Runtime-da pip/docker Hub asılılığı olmamalıdır")
    add_bullet(doc, "Session ~8 saat; güclü parol siyasəti")
    add_bullet(doc, "Loglar: docker compose logs; disk doldurmasın deyə rotasiya")
    add_bullet(doc, "Yeniləmə: yeni image USB → docker load → compose up (downtime pəncərəsi)")

    # ---- 11 ----
    add_heading(doc, "11. İş bölgüsü — kim nə edir", 1)
    add_table(
        doc,
        ["Rol", "Görüləcək işlər", "Çıxış / acceptance"],
        [
            [
                "NetAdmin",
                "Statik IP, DNS/hostname, firewall 8000/80/443, AxTrax 1433 çıxış, NTP, (SMTP)",
                "LAN-dan URL açılır; AxTrax SQL telnet/test OK",
            ],
            [
                "SysAdmin",
                "AlmaLinux quraşdırma, Docker, disk/NAS backup, chrony, SELinux, Nginx (opsional)",
                "docker ps yaşıl; backup cron işləyir",
            ],
            [
                "ERP ops",
                ".env doldurma, migrate, invite_portal_user, axtrax_sync_status, smoke testlər",
                "Portal/ERP/Security login + AxTrax event gəlir",
            ],
            [
                "Təhlükəsizlik / Reception",
                "Pilot smoke: qonaq, kart, təhlükəsizlik panel",
                "P0 xəta yox",
            ],
            [
                "Rəhbərlik",
                "Pilot şirkət razılığı, URL çapı, mail siyasəti qərarı",
                "DoD imzası",
            ],
        ],
    )

    # ---- 12 ----
    add_heading(doc, "12. Addım-addım layihə planı (xronoloji)", 1)
    add_table(
        doc,
        ["Mərhələ", "Müddət (təxmini)", "İşlər"],
        [
            ["A. Hazırlıq", "2–5 gün", "Resurs təsdiqi, məlumat cədvəli, offline image build"],
            ["B. OS + Docker", "1–2 gün", "AlmaLinux, firewall, chrony, Docker, SELinux yoxlama"],
            ["C. Deploy", "1 gün", ".env, compose up, migrate, rollar, ilk admin"],
            ["D. AxTrax", "0.5–1 gün", "1433, poller, sync_status, lazımsa backfill"],
            ["E. Backup", "0.5 gün", "cron/timer + bir restore drill"],
            ["F. Pilot smoke", "1–3 gün", "1–2 şirkət: invite, ticket, qonaq, Excel export"],
            ["G. Qəbul", "1 gün", "Offline uplink sınağı + imza"],
        ],
    )

    # ---- 13 ----
    add_heading(doc, "13. Qəbul / go-live checklist", 1)
    add_bullet(doc, "[ ] CP_ENV=production, SEED_DEMO=0, DEBUG=0")
    add_bullet(doc, "[ ] ALLOWED_HOSTS / CSRF origins LAN hostname və ya IP ilə düzgün")
    add_bullet(doc, "[ ] İnternet uplink kəsikdə belə LAN-dan portal/ERP açılır")
    add_bullet(doc, "[ ] Font/ikon CDN xətası yoxdur")
    add_bullet(doc, "[ ] Admin / Reception / Security / Resident login")
    add_bullet(doc, "[ ] Ticket yaratmaq + cavab")
    add_bullet(doc, "[ ] Qonaq pre-reg / reception")
    add_bullet(doc, "[ ] AxTrax event interval daxilində gəlir (axtrax_sync_status)")
    add_bullet(doc, "[ ] Excel giriş/çıxış export (portal + ERP CP işçilər + security)")
    add_bullet(doc, "[ ] Invite link (SMTP və ya əl ilə kopyalama)")
    add_bullet(doc, "[ ] Backup yazılır; restore drill sənədləşdirilib")
    add_bullet(doc, "[ ] docker restart / reboot sonrası xidmətlər qalxır (restart: unless-stopped)")

    # ---- 14 ----
    add_heading(doc, "14. Troubleshooting (qısa)", 1)
    add_table(
        doc,
        ["Problem", "Yoxlama", "Həll istiqaməti"],
        [
            ["Brauzer bağlana bilmir", "firewall-cmd, docker ps, port 8000", "firewall port; compose up"],
            ["CSRF / 403", "CSRF_TRUSTED_ORIGINS", "URL sxemi http/https uyğunluğu"],
            ["AxTrax empty", "ping/telnet 1433, password env", "şəbəkə + RO hesab; poller log"],
            ["DB connection", "POSTGRES_HOST=db, compose network", ".env və depends_on"],
            ["Permission denied media", "SELinux / volume rights", "kontekst və ya mount flags"],
            ["Demo data prod-da", "SEED_DEMO", "0 edin; wipe qərarı"],
        ],
    )

    # ---- 15 ----
    add_heading(doc, "15. Disk və kataloq strukturu", 1)
    add_code(
        doc,
        "/opt/citypoint/                 # repo (kod + bin/prod)\n"
        "/opt/citypoint/bin/prod/.env    # secret-lər (chmod 600)\n"
        "/opt/citypoint-images/          # docker save/load tar faylları\n"
        "/backup/citypoint/              # pg_dump + media arxivləri\n"
        "# Docker volumes (docker volume ls):\n"
        "#   ..._postgres_data\n"
        "#   ..._media\n"
        "#   ..._staticfiles",
    )
    add_para(
        doc,
        "Disk planı: OS + Docker (~40–60 GB), DB/media böyüməsi üçün ayrı volume və ya "
        "böyük /var/lib/docker, backup üçün /backup və ya NAS mount. "
        "df -h və docker system df həftəlik yoxlanılsın.",
    )

    # ---- 16 ----
    add_heading(doc, "16. NetAdmin ətraflı checklist", 1)
    add_bullet(doc, "[ ] ERP hosta statik IP verildi; gateway/DNS LAN-daxili")
    add_bullet(doc, "[ ] DNS və ya hosts: erp.citypoint.local / portal.citypoint.local → eyni IP")
    add_bullet(doc, "[ ] VLAN: staff ERP və resident Portal ayrımı (mümkünsə)")
    add_bullet(doc, "[ ] Firewall IN: yalnız LAN → 8000 (və ya 80/443); 5432 XARİCİDƏN bağlı")
    add_bullet(doc, "[ ] Firewall OUT: host → AxTrax SQL :1433; (opsional) daxili SMTP")
    add_bullet(doc, "[ ] Firewall OUT: runtime üçün Docker Hub / public SMTP tələb olunmur")
    add_bullet(doc, "[ ] AxTrax RO hesabı yalnız SELECT; yazma hüququ yox")
    add_bullet(doc, "[ ] NTP: daxili DC və ya chrony pool")
    add_bullet(doc, "[ ] Test: işçi PC-dən http://erp...:8000/login/ açılır")
    add_bullet(doc, "[ ] Test: ERP hostdan nc -vz <axtrax_ip> 1433 uğurlu")
    add_bullet(doc, "[ ] İnternet uplink kəsik sınağı (air-gap smoke)")

    # ---- 17 ----
    add_heading(doc, "17. Pilot smoke testləri (addım-addım)", 1)
    add_heading(doc, "17.1. Admin / ERP", 2)
    add_bullet(doc, "Admin ilə daxil ol → şirkətlər / istifadəçilər görünür")
    add_bullet(doc, "Reception: qonaq siyahısı, check-in/out axını")
    add_bullet(doc, "Tickets: yeni ticket + cavab")
    add_bullet(doc, "CP daxili işçilər: Access Excel export (şirkət adı faylda, sərhədlər)")
    add_heading(doc, "17.2. Resident Portal", 2)
    add_bullet(doc, "invite_portal_user ilə 1–2 istifadəçi")
    add_bullet(doc, "Invite link ilə parol təyin et (SMTP və ya əl ilə)")
    add_bullet(doc, "Portal: qonaq əlavə et, expected arrival 24h")
    add_bullet(doc, "Portal: Access Excel export")
    add_heading(doc, "17.3. Security", 2)
    add_bullet(doc, "/security/ login")
    add_bullet(doc, "Şirkət işçiləri + Access export")
    add_bullet(doc, "Rapid swipe / shaft alert (əgər F#Shaft* oxuyucular aktivdirsə)")
    add_heading(doc, "17.4. AxTrax", 2)
    add_bullet(doc, "axtrax_sync_status — son sync vaxtı yaxındır")
    add_bullet(doc, "Kart oxutma → AccessEvent-də görünür (bir neçə poll intervalı)")
    add_bullet(doc, "Lazım gələrsə backfill_axtrax_events --from=YYYY-MM-DD")

    # ---- 18 ----
    add_heading(doc, "18. Yeniləmə və rollback", 1)
    add_para(doc, "Yeniləmə (offline):", bold=True)
    add_bullet(doc, "Build maşınında yeni web image build → docker save")
    add_bullet(doc, "USB/NAS ilə serverə köçür → docker load")
    add_bullet(doc, "Backup götür (DB + media)")
    add_bullet(doc, "docker compose ... up -d (downtime pəncərəsi razılaşdırılsın)")
    add_bullet(doc, "migrate (entrypoint və ya əl ilə) + smoke")
    add_para(doc, "Rollback:", bold=True)
    add_bullet(doc, "Əvvəlki image tag-ini compose-da qaytarın və up -d")
    add_bullet(doc, "Əgər migrate geriyə uyğunsuzdursa: DB dump-dan bərpa (diqqət: data itkisi riski)")
    add_bullet(doc, "Hər major yeniləmədən əvvəl restore drill")

    # ---- 19 ----
    add_heading(doc, "19. Monitorinq və gündəlik ops", 1)
    add_bullet(doc, "docker compose ps — web/db Up")
    add_bullet(doc, "docker compose logs --tail=200 web — xəta axtarışı")
    add_bullet(doc, "axtrax_sync_status — poller geridə qalmır")
    add_bullet(doc, "df -h /backup və Docker disk")
    add_bullet(doc, "Son backup faylının tarixi / ölçüsü")
    add_bullet(doc, "Reboot test (ayda bir): systemctl reboot → compose auto-start")

    # ---- 20 ----
    add_heading(doc, "20. Əlavə A — .env nümunə sahələri", 1)
    add_code(
        doc,
        "CP_ENV=production\n"
        "CP_OFFLINE=1\n"
        "DJANGO_SECRET_KEY=<güçlü-təsadüfi>\n"
        "DJANGO_DEBUG=0\n"
        "DJANGO_ALLOWED_HOSTS=erp.citypoint.local,10.x.x.x\n"
        "DJANGO_CSRF_TRUSTED_ORIGINS=http://erp.citypoint.local\n"
        "POSTGRES_DB=citypoint\n"
        "POSTGRES_USER=cp_app\n"
        "POSTGRES_PASSWORD=<secret>\n"
        "POSTGRES_HOST=db\n"
        "POSTGRES_PORT=5432\n"
        "SEED_DEMO=0\n"
        "DEFAULT_FROM_EMAIL=noreply@citypoint.local\n"
        "EMAIL_FILE_PATH=/app/var/mail_outbox\n"
        "SESSION_COOKIE_AGE=28800\n"
        "AXTRAX_POLL_ON_START=1\n"
        "AXTRAX_POLL_INTERVAL=15\n"
        "TURNSTILE_MSSQL_HOST=172.31.104.10\n"
        "TURNSTILE_MSSQL_PORT=1433\n"
        "TURNSTILE_MSSQL_DATABASE=AxTrax1\n"
        "TURNSTILE_MSSQL_USER=ReadOnlyerp\n"
        "TURNSTILE_MSSQL_PASSWORD=<secret>",
    )
    add_para(doc, "Tam şablon: bin/prod/.env.example — real .env git-ə commit edilməsin.")

    # ---- 21 ----
    add_heading(doc, "21. Əlavə B — Qəbul imza forması", 1)
    add_table(
        doc,
        ["Rol", "Ad / soyad", "Tarix", "İmza"],
        [
            ["NetAdmin", "", "", ""],
            ["SysAdmin", "", "", ""],
            ["ERP ops", "", "", ""],
            ["Reception / Security", "", "", ""],
            ["Rəhbərlik", "", "", ""],
        ],
    )
    add_para(
        doc,
        "İmza: go-live checklist tamamlanıb; LAN air-gap smoke keçib; backup/restore drill qeydə alınıb.",
    )

    # ---- 22 ----
    add_heading(doc, "22. Əlaqəli sənədlər (repo)", 1)
    add_bullet(doc, "docs/deployment/CityPoint_AlmaLinux10_Deploy_Rehberi.docx (bu fayl)")
    add_bullet(doc, "docs/deployment/air-gapped-lan.md")
    add_bullet(doc, "docs/deployment/environments.md")
    add_bullet(doc, "docs/deployment/phase1-pilot-checklist.md")
    add_bullet(doc, "docs/deployment/offline-smoke-checklist.md")
    add_bullet(doc, "docs/security/network-checklist.md")
    add_bullet(doc, "docs/security/postgres-roles.md")
    add_bullet(doc, "docs/backup/phase1-backup-restore.md")
    add_bullet(doc, "docs/integrations/axtrax-poller-runbook.md")
    add_bullet(doc, "bin/prod/RUN.txt")

    add_heading(doc, "23. Nəticə", 1)
    add_para(
        doc,
        "City Point ERP/Portal AlmaLinux 10-da Docker Compose ilə LAN modelində işləyə bilər. "
        "Uğur üçün üç şərt kritikdir: (1) düzgün ölçülmüş host + Docker, "
        "(2) AxTrax və brauzer LAN şəbəkəsi, (3) backup/restore və pilot smoke tamamlanması. "
        "İnternet yalnız image/build mərhələsində lazımdır; istismar air-gap ola bilər.",
    )

    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("— City Point ERP · Daxili istifadə · Confidential —")
    set_run_font(r, size=9, color=RGBColor(0x64, 0x74, 0x8B))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
