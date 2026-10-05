from datetime import datetime, timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import Role
from apps.residents.access_export import access_xlsx_filename, build_access_xlsx, iter_company_attendance_rows
from apps.residents.models import AccessEvent, AccessEventType, ResidentCompany, ResidentEmployee

User = get_user_model()


class AccessExportServiceTests(TestCase):
    def setUp(self):
        self.company = ResidentCompany.objects.create(
            name="LB Audit", slug="lb-audit", portal_active=True, is_internal=False
        )
        self.other = ResidentCompany.objects.create(
            name="Other Co", slug="other-co", portal_active=True, is_internal=False
        )
        self.emp = ResidentEmployee.objects.create(
            company=self.company,
            full_name="Anar Abdullayev",
            card_number="12345",
            is_active=True,
        )
        other_emp = ResidentEmployee.objects.create(
            company=self.other,
            full_name="Other Emp",
            card_number="999",
            is_active=True,
        )
        day = timezone.localdate().replace(day=1)
        in_at = timezone.make_aware(datetime.combine(day, datetime.strptime("09:15", "%H:%M").time()))
        out_at = timezone.make_aware(datetime.combine(day, datetime.strptime("18:40", "%H:%M").time()))
        AccessEvent.objects.create(
            employee=self.emp,
            event_type=AccessEventType.IN,
            occurred_at=in_at,
            employee_name=self.emp.full_name,
            card_number=self.emp.card_number,
        )
        AccessEvent.objects.create(
            employee=self.emp,
            event_type=AccessEventType.OUT,
            occurred_at=out_at,
            employee_name=self.emp.full_name,
            card_number=self.emp.card_number,
        )
        AccessEvent.objects.create(
            employee=other_emp,
            event_type=AccessEventType.IN,
            occurred_at=in_at,
            employee_name=other_emp.full_name,
            card_number=other_emp.card_number,
        )
        self.day = day
        self.date_to = day.replace(day=min(3, 28))

    def test_rows_scoped_to_company(self):
        rows = iter_company_attendance_rows(self.company, self.day, self.day)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["employee_name"], "Anar Abdullayev")
        self.assertEqual(rows[0]["card_number"], "12345")
        self.assertIsNotNone(rows[0]["in_at"])
        self.assertIsNotNone(rows[0]["out_at"])

    def test_xlsx_citypoint_report_layout_no_qeyd(self):
        payload = build_access_xlsx(self.company, self.day, self.date_to)
        wb = load_workbook(BytesIO(payload))
        ws = wb.active
        self.assertEqual(ws.title, "Hesabat")
        self.assertEqual(ws["A1"].value, "Şöbə: LB Audit")
        self.assertTrue(str(ws["A2"].value).startswith("İşçi:"))
        self.assertIn("Abdullayev,Anar", ws["A2"].value)
        headers = [ws.cell(3, c).value for c in range(1, 6)]
        self.assertEqual(headers, ["Tarix", "Gün", "Giriş", "Çıxış", "Müddət"])
        self.assertIsNone(ws.cell(3, 6).value)
        self.assertEqual(ws.cell(4, 1).value, self.day.strftime("%d.%m.%Y"))
        self.assertEqual(ws.cell(4, 3).value, "09:15")
        self.assertEqual(ws.cell(4, 4).value, "18:40")
        self.assertEqual(ws.cell(4, 5).value, "9:25")
        # Day without punch uses placeholder, no Qeyd column
        if self.date_to > self.day:
            self.assertEqual(ws.cell(5, 3).value, "----------")
            self.assertEqual(ws.cell(5, 4).value, "----------")
            self.assertEqual(ws.cell(5, 5).value, "0:00")
        self.assertEqual(
            access_xlsx_filename(self.company, self.day, self.date_to),
            f"LB_Audit_Report_{['','January','February','March','April','May','June','July','August','September','October','November','December'][self.day.month]}_AZ.xlsx",
        )


class AccessExportViewTests(TestCase):
    def setUp(self):
        self.company = ResidentCompany.objects.create(
            name="LB Audit", slug="lb-audit", portal_active=True, is_internal=False
        )
        self.other = ResidentCompany.objects.create(
            name="Other Co", slug="other-co", portal_active=True, is_internal=False
        )
        self.emp = ResidentEmployee.objects.create(
            company=self.company, full_name="Ali Veli", card_number="12345", is_active=True
        )
        AccessEvent.objects.create(
            employee=self.emp,
            event_type=AccessEventType.IN,
            occurred_at=timezone.now(),
            employee_name=self.emp.full_name,
            card_number=self.emp.card_number,
        )
        self.resident = User.objects.create_user(
            username="res-lb",
            email="res-lb@test.az",
            password="Pass12345!",
            role=Role.RESIDENT_USER,
            resident_company=self.company,
        )
        self.other_user = User.objects.create_user(
            username="res-other",
            email="res-other@test.az",
            password="Pass12345!",
            role=Role.RESIDENT_USER,
            resident_company=self.other,
        )
        self.admin = User.objects.create_user(
            username="admin-exp",
            email="admin-exp@test.az",
            password="Pass12345!",
            role=Role.ADMIN,
        )
        self.client = Client()

    def test_portal_export_own_company(self):
        self.client.force_login(self.resident)
        url = reverse("portal:employees_access_export")
        resp = self.client.get(url, {"range": "month"})
        self.assertEqual(resp.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resp["Content-Type"],
        )
        self.assertIn("LB_Audit_Report_", resp["Content-Disposition"])
        self.assertTrue(resp["Content-Disposition"].endswith('_AZ.xlsx"'))

    def test_portal_cannot_export_other_employee(self):
        foreign = ResidentEmployee.objects.create(
            company=self.other, full_name="X", card_number="1", is_active=True
        )
        self.client.force_login(self.resident)
        resp = self.client.get(reverse("portal:employee_access_export", args=[foreign.pk]))
        self.assertEqual(resp.status_code, 404)

    def test_erp_admin_export(self):
        self.client.force_login(self.admin)
        url = reverse("erp:resident_access_export", args=[self.company.slug])
        resp = self.client.get(
            url,
            {
                "range": "custom",
                "from": (timezone.localdate() - timedelta(days=1)).isoformat(),
                "to": timezone.localdate().isoformat(),
            },
        )
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.content))
        self.assertEqual(wb.active["A1"].value, "Şöbə: LB Audit")
        self.assertEqual(wb.active.cell(3, 1).value, "Tarix")

    def test_other_resident_blocked_from_erp(self):
        self.client.force_login(self.other_user)
        url = reverse("erp:resident_access_export", args=[self.company.slug])
        resp = self.client.get(url, {"range": "month"})
        self.assertIn(resp.status_code, (302, 403))


class InternalAndSecurityAccessExportTests(TestCase):
    def setUp(self):
        self.internal = ResidentCompany.objects.create(
            name="City Point", slug="city-point", portal_active=False, is_internal=True
        )
        self.emp = ResidentEmployee.objects.create(
            company=self.internal, full_name="CP Staff", card_number="100", is_active=True
        )
        AccessEvent.objects.create(
            employee=self.emp,
            event_type=AccessEventType.IN,
            occurred_at=timezone.now(),
            employee_name=self.emp.full_name,
            card_number=self.emp.card_number,
        )
        self.admin = User.objects.create_user(
            username="admin-cp",
            email="admin-cp@test.az",
            password="Pass12345!",
            role=Role.ADMIN,
        )
        self.security = User.objects.create_user(
            username="sec-cp",
            email="sec-cp@test.az",
            password="Pass12345!",
            role=Role.SECURITY,
        )
        self.client = Client()

    def test_erp_internal_staff_export(self):
        self.client.force_login(self.admin)
        resp = self.client.get(
            reverse("erp:internal_staff_access_export"),
            {"range": "month"},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertIn("City_Point_Report_", resp["Content-Disposition"])

    def test_security_city_point_export(self):
        self.client.force_login(self.security)
        resp = self.client.get(
            reverse("security:company_access_export", args=[self.internal.pk]),
            {"range": "month"},
        )
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.content))
        self.assertEqual(wb.active["A1"].value, "Şöbə: City Point")
