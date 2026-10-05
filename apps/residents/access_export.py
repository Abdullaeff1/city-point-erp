"""CityPoint_Report_* style attendance Excel (per employee calendar grid, no Qeyd column)."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO

from django.contrib.contenttypes.models import ContentType
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from apps.residents.access import attendance_for_day, employee_roster_qs
from apps.residents.models import AccessEvent, ResidentEmployee

AZ_WEEKDAYS = (
    "Bazar ertəsi",
    "Çərşənbə axşamı",
    "Çərşənbə",
    "Cümə axşamı",
    "Cümə",
    "Şənbə",
    "Bazar",
)

MONTHS_EN = (
    "",
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)

MISSING = "----------"


def _daterange(date_from: date, date_to: date):
    cur = date_from
    while cur <= date_to:
        yield cur
        cur += timedelta(days=1)


def _fmt_time(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%H:%M")


def _fmt_duration(in_at, out_at) -> str:
    if not in_at or not out_at:
        return "0:00"
    start = timezone.localtime(in_at)
    end = timezone.localtime(out_at)
    seconds = int((end - start).total_seconds())
    if seconds < 0:
        return "0:00"
    hours, rem = divmod(seconds, 3600)
    minutes = rem // 60
    return f"{hours}:{minutes:02d}"


def _split_name(full_name: str) -> tuple[str, str]:
    parts = (full_name or "").strip().split()
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    # AxTrax style: Last,First[ Middle]
    return parts[-1], " ".join(parts[:-1])


def _axtrax_id_map(employee_ids: list[int]) -> dict[int, int]:
    """ResidentEmployee.pk → AxTrax IdEmpNum via ExternalIdentity emp:N."""
    if not employee_ids:
        return {}
    from apps.integrations.axtrax_people_sync import AXTRAX_SYSTEM
    from apps.integrations.models import ExternalIdentity

    ct = ContentType.objects.get_for_model(ResidentEmployee)
    mapping: dict[int, int] = {}
    for ext_id, entity_id in ExternalIdentity.objects.filter(
        system=AXTRAX_SYSTEM,
        entity_type=ct,
        entity_id__in=employee_ids,
        external_id__startswith="emp:",
    ).values_list("external_id", "entity_id"):
        try:
            mapping[int(entity_id)] = int(str(ext_id).split(":", 1)[1])
        except (IndexError, ValueError, TypeError):
            continue
    return mapping


def _employee_heading(emp: ResidentEmployee, axtrax_id: int | None) -> str:
    last, first = _split_name(emp.full_name)
    if last and first:
        name = f"{last},{first}"
    else:
        name = emp.full_name or ""
    card = (emp.card_number or "").strip()
    bits = ["İşçi:", name]
    if card:
        bits.append(card)
    if axtrax_id:
        bits.append(f"#{axtrax_id}")
    return " ".join(bits)


def load_company_day_attendance(
    company,
    date_from,
    date_to,
    *,
    employee: ResidentEmployee | None = None,
    roster: str = "active",
) -> tuple[list[ResidentEmployee], dict[int, dict[date, tuple]]]:
    """Employees + map emp_id → day → (in_at, out_at, status)."""
    qs = ResidentEmployee.objects.filter(company=company)
    if employee is not None:
        qs = qs.filter(pk=employee.pk)
    else:
        qs, _ = employee_roster_qs(qs, roster)
    employees = list(qs.order_by("full_name"))
    by_emp_day: dict[int, dict[date, tuple]] = defaultdict(dict)
    if not employees:
        return employees, by_emp_day

    emp_ids = [e.pk for e in employees]
    start_dt = timezone.make_aware(datetime.combine(date_from, datetime.min.time()))
    end_dt = timezone.make_aware(datetime.combine(date_to + timedelta(days=1), datetime.min.time()))
    events = AccessEvent.objects.filter(
        employee_id__in=emp_ids,
        occurred_at__gte=start_dt,
        occurred_at__lt=end_dt,
    ).order_by("occurred_at", "id")

    grouped: dict[int, dict[date, list]] = defaultdict(lambda: defaultdict(list))
    for event in events:
        day = timezone.localtime(event.occurred_at).date()
        grouped[event.employee_id][day].append(event)

    for emp_id, days in grouped.items():
        for day, day_events in days.items():
            by_emp_day[emp_id][day] = attendance_for_day(day_events)
    return employees, by_emp_day


# Back-compat for older imports/tests
def iter_company_attendance_rows(company, date_from, date_to, *, employee=None, roster="active"):
    employees, by_emp_day = load_company_day_attendance(
        company, date_from, date_to, employee=employee, roster=roster
    )
    rows = []
    for emp in employees:
        for day in sorted(by_emp_day[emp.pk].keys()):
            in_at, out_at, status = by_emp_day[emp.pk][day]
            rows.append(
                {
                    "date": day,
                    "employee": emp,
                    "employee_name": emp.full_name,
                    "card_number": emp.card_number or "",
                    "in_at": in_at,
                    "out_at": out_at,
                    "status": status,
                }
            )
    rows.sort(key=lambda r: (r["date"], (r["employee_name"] or "").casefold()))
    return rows


def _safe_filename_part(value: str) -> str:
    text = (value or "").strip()
    out = []
    for ch in text:
        if ch.isalnum() or ch in ("-", "_"):
            out.append(ch)
        elif ch.isspace() or ch in (".", "/", "\\", ","):
            out.append("_")
    cleaned = "".join(out).strip("_")
    while "__" in cleaned:
        cleaned = cleaned.replace("__", "_")
    return cleaned or "Company"


def access_xlsx_filename(company, date_from, date_to, *, employee=None) -> str:
    company_part = _safe_filename_part(getattr(company, "name", None) or getattr(company, "slug", "") or "Company")
    if date_from.month == date_to.month and date_from.year == date_to.year:
        month = MONTHS_EN[date_from.month]
        return f"{company_part}_Report_{month}_AZ.xlsx"
    return f"{company_part}_Report_{date_from.isoformat()}_{date_to.isoformat()}_AZ.xlsx"


def build_access_xlsx(
    company,
    date_from,
    date_to,
    *,
    employee: ResidentEmployee | None = None,
    roster: str = "active",
) -> bytes:
    """Match CityPoint_Report_September_AZ layout without the Qeyd column."""
    employees, by_emp_day = load_company_day_attendance(
        company, date_from, date_to, employee=employee, roster=roster
    )
    ax_ids = _axtrax_id_map([e.pk for e in employees])

    wb = Workbook()
    ws = wb.active
    ws.title = "Hesabat"

    title_font = Font(name="Calibri", size=11, bold=True)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=11)
    left = Alignment(horizontal="left", vertical="center")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="64748B")
    medium = Side(style="medium", color="1F497D")
    emp_fill = PatternFill("solid", fgColor="D6E3F0")
    header_fill = PatternFill("solid", fgColor="1F497D")
    alt_fill = PatternFill("solid", fgColor="F8FAFC")

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 12

    def _box_border(*, left_edge=False, right_edge=False, top_edge=False, bottom_edge=False):
        return Border(
            left=medium if left_edge else thin,
            right=medium if right_edge else thin,
            top=medium if top_edge else thin,
            bottom=medium if bottom_edge else thin,
        )

    row = 1
    # Department once at top (sample report)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    dept = ws.cell(row=row, column=1, value=f"Şöbə: {company.name}")
    dept.font = title_font
    dept.alignment = left
    for col in range(1, 6):
        cell = ws.cell(row=row, column=col)
        cell.fill = emp_fill
        cell.border = _box_border(left_edge=col == 1, right_edge=col == 5, top_edge=True)
    row += 1

    headers = ("Tarix", "Gün", "Giriş", "Çıxış", "Müddət")
    days = list(_daterange(date_from, date_to))

    for idx, emp in enumerate(employees):
        if idx > 0:
            row += 1  # blank separator between employees

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        emp_cell = ws.cell(
            row=row,
            column=1,
            value=_employee_heading(emp, ax_ids.get(emp.pk)),
        )
        emp_cell.font = title_font
        emp_cell.alignment = left
        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            cell.fill = emp_fill
            cell.border = _box_border(left_edge=col == 1, right_edge=col == 5, top_edge=True)
        row += 1

        for col, label in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = _box_border(left_edge=col == 1, right_edge=col == 5)
        row += 1

        day_map = by_emp_day.get(emp.pk, {})
        for day_i, day in enumerate(days):
            punch = day_map.get(day)
            if punch:
                in_at, out_at, _status = punch
                in_s = _fmt_time(in_at) or MISSING
                out_s = _fmt_time(out_at) if out_at else ("" if in_at else MISSING)
                if in_at and not out_at:
                    out_s = ""
                elif not in_at and not out_at:
                    in_s = out_s = MISSING
                dur = _fmt_duration(in_at, out_at)
            else:
                in_s = out_s = MISSING
                dur = "0:00"

            values = [
                day.strftime("%d.%m.%Y"),
                AZ_WEEKDAYS[day.weekday()],
                in_s,
                out_s,
                dur,
            ]
            is_last = day_i == len(days) - 1
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = cell_font
                cell.alignment = center if col != 2 else left
                if day_i % 2 == 1:
                    cell.fill = alt_fill
                cell.border = _box_border(
                    left_edge=col == 1,
                    right_edge=col == 5,
                    bottom_edge=is_last,
                )
            row += 1

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def access_xlsx_http_response(
    company,
    date_from,
    date_to,
    *,
    employee: ResidentEmployee | None = None,
    roster: str = "active",
) -> HttpResponse:
    payload = build_access_xlsx(
        company,
        date_from,
        date_to,
        employee=employee,
        roster=roster,
    )
    filename = access_xlsx_filename(company, date_from, date_to, employee=employee)
    response = HttpResponse(
        payload,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
