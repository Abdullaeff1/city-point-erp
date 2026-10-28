"""Import City Point lease / vacant spaces from the commercial Excel workbook."""

from __future__ import annotations

import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.db import transaction
from django.utils.text import slugify

from apps.property.models import (
    Building,
    CommercialStatus,
    Floor,
    Occupancy,
    OperationalStatus,
    Space,
)
from apps.residents.models import ResidentCompany

# Excel status labels → (occupancy, commercial_status). Keys are matched via _fold().
_STATUS_RAW = {
    "İcarədə": (Occupancy.ACTIVE_LEASE, CommercialStatus.ACTIVE),
    "Boş sahə": (Occupancy.VACANT, CommercialStatus.VACANT),
    "Bron edilib": (Occupancy.VACANT, CommercialStatus.RESERVED),
    "Müvəqqəti rezervdə(Boş)": (Occupancy.VACANT, CommercialStatus.RESERVED),
    "Təmirsiz": (Occupancy.MAINTENANCE, CommercialStatus.VACANT),
    "Daxili istifadədə": (Occupancy.ACTIVE_LEASE, CommercialStatus.ACTIVE),
}

# Explicit Excel company label → existing ResidentCompany.name (or create name)
_COMPANY_ALIASES_RAW = {
    "Food point 13": "Food Point 13",
    "COFIESTO": "Cofıesto",
    "Psixoloq- Leyla Şamilova": "PSİXOLOQ",
    "CJ ICM FZCO (baş ofis)": "CJ ICM",
    "CJ ICM FZCO": "CJ ICM",
    "Spark Aİ": "Spark AI",
    "ATL TECH": "ATL Tech",
    "Epza&Quetera Park": "EPZA   CONSULTİNG  MMC",
    "Baltic transline": "Baltic Transline Azerbaij",
    "SOFTGİLE": "Softgile",
    "NİTEL TEKNOLOJİ": "NITEL",
    "ORDU İDMAN KLUBU": "ORDU İDMAN KLUBU İB",
    "ATB": "A.T.Bank",
    "LB AUDİT (INVEST FINANS)": "L.B.Audit",
    "ADMİNİSTRASİYA": "City Point",
    "Ryumka": "RYUMKA MMC",
    "Referans": "Referans",
    "BİH and Referans": "BIH PHARMACEUTİCALS LTD",
    "Tvins group": "Tvins",
    "Twins group Park Elmler": "PARK ELMLƏR",
    "MRT MMC": "MRT MMC",
    "AZİMUTH & GEOTRONICS": "Geotranik",
    "GNR Altyapı": "GNR ALTYAPI İNŞAAT ANONİM",
    "DATAPAY": "Datapay MMC",
    "AMK": "AMK Medical",
    "PMİS": "PMİS MMC",
    "SİMPLEX": "Simplex QSC",
    "ST Mark": "ST Mark",
    "Qala Express": "Qala Express",
    "NESTLE": "Nestle",
    "Prodata": "ProData",
    "Project 13": "Project 13",
    "JM DMC": "JM DMC",
    "Zirvion Bros": "ZİRVİON MMC",
    "Delina dizayn": "Delina",
    "Moon Group": "Moon Group",
    "NORAS": "NORAS MMC",
    "ASBC": "ASBC",
    "TİD LEGAL SOLUTION": "TID consulting",
    "BAKER TİLLY": "Bakertilly",
    "SİMBRELLA": "Simbrella",
    "INTEGRO CONSULTING": "İNTEGRO CONSULTİNG",
    "VOİD ARCİTECH": "VOİD ARCHİTECTURE",
    "BAKI FOND BİRJASI (1)": "Bakı Fond Birjası QSC",
    "BAKI FOND BİRJASI (2)": "Bakı Fond Birjası QSC",
    "BAKI FOND BİRJASI": "Bakı Fond Birjası QSC",
}


def _fold(text: str) -> str:
    text = unicodedata.normalize("NFKD", (text or "").strip())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("İ", "i").replace("I", "i").replace("ı", "i")
    text = text.casefold()
    text = text.replace("ə", "e").replace("ö", "o").replace("ü", "u")
    text = text.replace("ğ", "g").replace("ç", "c").replace("ş", "s")
    text = re.sub(r"[^\w\s&+-]+", " ", text, flags=re.UNICODE)
    text = re.sub(r"\s+", " ", text).strip()
    return text


STATUS_MAP = {_fold(k): v for k, v in _STATUS_RAW.items()}
COMPANY_ALIASES = {_fold(k): v for k, v in _COMPANY_ALIASES_RAW.items()}


def _norm_company_key(name: str) -> str:
    key = _fold(name)
    for noise in (
        " mmc",
        " qsc",
        " ltd",
        " fzco",
        " ib",
        " anonim",
        " sirketi",
        " şirkəti",
    ):
        key = key.replace(_fold(noise).strip(), "")
    return re.sub(r"\s+", " ", key).strip()


def floor_code_for(value) -> str | None:
    if value is None or value == "":
        return None
    try:
        n = int(float(value))
    except (TypeError, ValueError):
        return None
    if n < 0:
        return f"B{abs(n):02d}"
    return f"F{n:02d}"


def floor_name_for(code: str) -> str:
    if code.startswith("B"):
        return f"Yeraltı {int(code[1:])}"
    n = int(code[1:])
    return f"{n}-ci mərtəbə"


def _dec(value, default="0") -> Decimal:
    if value is None or value == "":
        return Decimal(default)
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        return Decimal(default)


def _split_combo(label: str):
    """Split 'A (60 m2) & B (20 m2)' into [(name, area|None), ...]."""
    text = (label or "").strip()
    if "&" not in text and " və " not in text.casefold():
        return [(text, None)]
    parts = re.split(r"\s*&\s*|\s+və\s+", text, flags=re.IGNORECASE)
    out = []
    for part in parts:
        part = part.strip(" ,")
        if not part:
            continue
        m = re.search(r"^(.*?)\s*\(\s*([\d.,]+)\s*m2\s*\)\s*$", part, flags=re.IGNORECASE)
        if m:
            area = _dec(m.group(2).replace(",", "."))
            out.append((m.group(1).strip(), area))
        else:
            out.append((part, None))
    return out or [(text, None)]


def parse_workbook(path: str | Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_name = next(
        (n for n in wb.sheetnames if "İcarə" in n or "Icarə" in n or "icar" in n.casefold()),
        wb.sheetnames[0],
    )
    ws = wb[sheet_name]
    rows = []
    for excel_row, raw in enumerate(ws.iter_rows(min_row=3, max_col=8, values_only=True), 3):
        num, company, status, floor, area, price, amount, notes = raw
        company = (str(company).strip() if company is not None else "")
        status_raw = (str(status).strip() if status is not None else "")
        if not company or company.upper().startswith("CƏMİ") or company.upper().startswith("CEMI"):
            continue
        status_key = _fold(status_raw)
        if status_key not in STATUS_MAP:
            continue
        occupancy, commercial = STATUS_MAP[status_key]
        fcode = floor_code_for(floor)
        if not fcode:
            continue
        total_area = _dec(area)
        price_m2 = _dec(price) if price not in (None, "") else None
        note = (str(notes).strip() if notes else "")
        try:
            seq = int(float(num))
        except (TypeError, ValueError):
            seq = excel_row

        parts = _split_combo(company)
        # Shared row without per-part areas → one space with full label.
        if len(parts) == 1 or all(a is None for _, a in parts):
            parts = [(company, total_area)]

        for idx, (part_name, part_area) in enumerate(parts, start=1):
            area_m2 = part_area if part_area is not None else total_area
            suffix = f"-{idx}" if len(parts) > 1 else ""
            code = f"CP-L{seq:03d}{suffix}"
            vacantish = _fold(part_name).startswith("bos")
            rows.append(
                {
                    "excel_row": excel_row,
                    "seq": seq,
                    "code": code,
                    "label": part_name.strip(),
                    "status_raw": status_raw,
                    "occupancy": occupancy,
                    "commercial_status": commercial,
                    "floor_code": fcode,
                    "area_m2": area_m2,
                    "base_rent_rate": price_m2 if price_m2 and price_m2 > 0 else None,
                    "notes": note,
                    "is_vacant_label": vacantish,
                    "fit_out_status": "Təmirsiz" if "temirsiz" in status_key or "təmirsiz" in status_key else "",
                }
            )
    return rows


def _company_index():
    by_norm = {}
    for c in ResidentCompany.objects.all():
        by_norm[_norm_company_key(c.name)] = c
        by_norm[_fold(c.name)] = c
    return by_norm


def resolve_company(label: str, index: dict, *, create_missing: bool) -> ResidentCompany | None:
    key = _fold(label)
    if key.startswith("bos"):
        return None
    alias_target = COMPANY_ALIASES.get(key)
    candidates = []
    if alias_target:
        candidates.append(alias_target)
    candidates.append(label)

    for cand in candidates:
        hit = index.get(_norm_company_key(cand)) or index.get(_fold(cand))
        if hit:
            return hit
        # partial contains
        nkey = _norm_company_key(cand)
        for k, company in index.items():
            if nkey and (nkey in k or k in nkey) and len(nkey) >= 4:
                return company

    if not create_missing:
        return None
    name = alias_target or label.strip()
    base = slugify(name) or f"company-{abs(hash(name)) % 100000}"
    slug = base[:80]
    i = 2
    while ResidentCompany.objects.filter(slug=slug).exists():
        slug = f"{base[:70]}-{i}"
        i += 1
    company = ResidentCompany.objects.create(
        name=name[:160],
        slug=slug,
        status="active",
        portal_active=False,
        is_internal=(name.strip().casefold() == "city point"),
    )
    index[_norm_company_key(company.name)] = company
    index[_fold(company.name)] = company
    return company


def ensure_floors(floor_codes: set[str]) -> dict[str, Floor]:
    building, _ = Building.objects.get_or_create(
        code="CP",
        defaults={"name": "City Point Business Center"},
    )
    out = {}
    for code in sorted(floor_codes):
        floor, _ = Floor.objects.get_or_create(
            building=building,
            code=code,
            defaults={"name": floor_name_for(code)},
        )
        out[code] = floor
    return out


def preview_import(path: str | Path) -> list[dict]:
    rows = parse_workbook(path)
    index = _company_index()
    preview = []
    for row in rows:
        company = None
        if not row["is_vacant_label"] and row["occupancy"] == Occupancy.ACTIVE_LEASE:
            company = resolve_company(row["label"], index, create_missing=False)
        elif row["label"].strip().upper().startswith("ADMİN") or _fold(row["label"]) == "administrasiya":
            company = resolve_company("ADMİNİSTRASİYA", index, create_missing=False)
        elif row["commercial_status"] == CommercialStatus.RESERVED and not row["is_vacant_label"]:
            company = resolve_company(row["label"], index, create_missing=False)
        preview.append({**row, "company": company.name if company else None, "company_id": company.pk if company else None})
    return preview


@transaction.atomic
def apply_import(path: str | Path, *, create_missing_companies: bool = True) -> dict:
    rows = parse_workbook(path)
    floors = ensure_floors({r["floor_code"] for r in rows})
    index = _company_index()
    # Prefer City Point for admin spaces.
    city = ResidentCompany.objects.filter(slug="city-point").first()
    if city:
        index[_fold("city point")] = city
        index[_norm_company_key("City Point")] = city

    created = updated = 0
    linked = unmatched = 0
    created_companies = []

    # Drop the demo ASBC mismatch so Excel ASBC row owns the link cleanly.
    demo = Space.objects.filter(code="CP-F08-OFF-801").select_related("resident").first()
    if demo and demo.resident and demo.resident.slug == "asbc":
        demo.resident = None
        demo.save(update_fields=["resident"])

    for row in rows:
        company = None
        label_fold = _fold(row["label"])
        if label_fold == "administrasiya" or label_fold.startswith("administrasiya"):
            company = city or resolve_company("ADMİNİSTRASİYA", index, create_missing=create_missing_companies)
        elif not row["is_vacant_label"]:
            before = set(ResidentCompany.objects.values_list("pk", flat=True))
            should_create = create_missing_companies and (
                row["occupancy"] == Occupancy.ACTIVE_LEASE
                or (
                    row["commercial_status"] == CommercialStatus.RESERVED
                    and not row["is_vacant_label"]
                )
            )
            company = resolve_company(row["label"], index, create_missing=should_create)
            after = set(ResidentCompany.objects.values_list("pk", flat=True))
            for pk in after - before:
                created_companies.append(ResidentCompany.objects.get(pk=pk).name)

        if company:
            linked += 1
        elif not row["is_vacant_label"] and row["occupancy"] == Occupancy.ACTIVE_LEASE:
            unmatched += 1

        name = row["label"][:120]
        defaults = {
            "name": name,
            "floor": floors[row["floor_code"]],
            "resident": company,
            "area_m2": row["area_m2"],
            "rentable_area_m2": row["area_m2"],
            "occupancy": row["occupancy"],
            "commercial_status": row["commercial_status"],
            "operational_status": (
                OperationalStatus.MAINTENANCE
                if row["occupancy"] == Occupancy.MAINTENANCE
                else OperationalStatus.AVAILABLE
            ),
            "base_rent_rate": row["base_rent_rate"],
            "fit_out_status": row["fit_out_status"],
            "public_description": row["notes"],
            "is_public": row["occupancy"] == Occupancy.VACANT
            and row["commercial_status"] in {CommercialStatus.VACANT, CommercialStatus.RESERVED},
        }
        obj, was_created = Space.objects.update_or_create(code=row["code"], defaults=defaults)
        if was_created:
            created += 1
        else:
            updated += 1

    return {
        "rows": len(rows),
        "created": created,
        "updated": updated,
        "linked": linked,
        "unmatched_leased": unmatched,
        "created_companies": created_companies,
        "floors": sorted(floors),
    }
