"""Read a building-asset list from Excel or Word and save it."""

from __future__ import annotations

import io
from datetime import date, datetime

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook

from apps.property.models import BuildingAsset, BuildingAssetCategory, Floor, add_one_month

SESSION_KEY = "building_asset_import"
MAX_ROWS = 500

_HEADER_ALIASES = {
    "ad": "name",
    "adi": "name",
    "name": "name",
    "avadanliq": "name",
    "cihaz": "name",
    "kod": "code",
    "code": "code",
    "kateqoriya": "category",
    "category": "category",
    "nov": "category",
    "mertebe": "floor",
    "floor": "floor",
    "yer": "location",
    "location": "location",
    "say": "quantity",
    "quantity": "quantity",
    "miqdar": "quantity",
    "son baxis": "last_inspected_on",
    "baxis": "last_inspected_on",
    "qeyd": "notes",
    "notes": "notes",
}

_CATEGORY_ALIASES = {
    "hvac": BuildingAssetCategory.HVAC,
    "lift": BuildingAssetCategory.LIFT,
    "elektrik": BuildingAssetCategory.ELECTRIC,
    "electric": BuildingAssetCategory.ELECTRIC,
    "yangin": BuildingAssetCategory.FIRE,
    "fire": BuildingAssetCategory.FIRE,
    "nasos": BuildingAssetCategory.PUMP,
    "pump": BuildingAssetCategory.PUMP,
    "mebel": BuildingAssetCategory.FURNITURE,
    "furniture": BuildingAssetCategory.FURNITURE,
    "diger": BuildingAssetCategory.OTHER,
    "other": BuildingAssetCategory.OTHER,
}


def _norm(value) -> str:
    text = str(value or "").strip().replace("İ", "i").replace("I", "i")
    text = text.casefold()
    for src, dst in (
        ("ı", "i"),
        ("ə", "e"),
        ("ö", "o"),
        ("ü", "u"),
        ("ş", "s"),
        ("ç", "c"),
        ("ğ", "g"),
    ):
        text = text.replace(src, dst)
    for char in "/,;:|-":
        text = text.replace(char, " ")
    return " ".join(text.split())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _header_field(header) -> str:
    text = _norm(header)
    if not text:
        return ""
    exact = _HEADER_ALIASES.get(text)
    if exact:
        return exact
    if "son" in text and "baxis" in text:
        return "last_inspected_on"
    return ""


def _map_headers(headers) -> dict[int, str]:
    mapped = {}
    for index, header in enumerate(headers):
        key = _header_field(header)
        if key and key not in mapped.values():
            mapped[index] = key
    return mapped


def _category_code(value) -> str:
    text = _norm(value)
    if not text:
        return BuildingAssetCategory.OTHER
    if text in _CATEGORY_ALIASES:
        return _CATEGORY_ALIASES[text]
    for alias, code in _CATEGORY_ALIASES.items():
        if alias in text:
            return code
    return BuildingAssetCategory.OTHER


def _quantity(value) -> int:
    text = _cell_text(value)
    if not text:
        return 1
    try:
        return max(1, int(float(text.replace(",", "."))))
    except ValueError:
        return 1


def _floor_number(text: str):
    norm = _norm(text).replace(" ", "")
    if norm.startswith("f") and norm[1:].isdigit():
        return int(norm[1:])
    digits = ""
    for char in norm:
        if char.isdigit():
            digits += char
        elif digits:
            break
    return int(digits) if digits else None


def match_floor(raw, floors):
    text = _cell_text(raw)
    if not text:
        return None
    wanted = _norm(text)
    wanted_number = _floor_number(text)
    numbered = []
    for floor in floors:
        if _norm(floor.code) == wanted or _norm(floor.name) == wanted:
            return floor
        number = _floor_number(floor.code) or _floor_number(floor.name)
        if wanted_number is not None and number == wanted_number:
            numbered.append(floor)
    if len(numbered) == 1:
        return numbered[0]
    return None


def _row_from_mapping(raw, mapping) -> dict:
    item = {
        "name": "",
        "code": "",
        "category": "",
        "floor": "",
        "location": "",
        "quantity": 1,
        "notes": "",
        "last_inspected_on": None,
    }
    for index, key in mapping.items():
        if index >= len(raw):
            continue
        value = raw[index]
        if key == "quantity":
            item["quantity"] = _quantity(value)
        elif key == "last_inspected_on":
            item["last_inspected_on"] = _parse_date(value)
        else:
            item[key] = _cell_text(value)
    return item


def _rows_from_table(table_rows: list[list]) -> list[dict]:
    if not table_rows:
        return []
    mapping = _map_headers(table_rows[0])
    if "name" in mapping.values():
        data_rows = table_rows[1:]
    else:
        mapping = {0: "name"}
        data_rows = table_rows
    parsed = []
    for raw in data_rows:
        if not any(_cell_text(cell) for cell in raw):
            continue
        item = _row_from_mapping(raw, mapping)
        if not item["name"]:
            continue
        parsed.append(item)
        if len(parsed) > MAX_ROWS:
            raise ValueError(f"Siyahı {MAX_ROWS} sətirdən çox ola bilməz.")
    return parsed


def parse_asset_xlsx(file_obj) -> list[dict]:
    workbook = load_workbook(file_obj, data_only=True)
    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        return _rows_from_table(rows)
    finally:
        workbook.close()


def parse_asset_docx(file_obj) -> list[dict]:
    from docx import Document

    document = Document(file_obj)
    if document.tables:
        table = document.tables[0]
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        return _rows_from_table(rows)
    parsed = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        parsed.append(
            {
                "name": text[:160],
                "code": "",
                "category": "",
                "floor": "",
                "location": "",
                "quantity": 1,
                "notes": "",
                "last_inspected_on": None,
            }
        )
        if len(parsed) > MAX_ROWS:
            raise ValueError(f"Siyahı {MAX_ROWS} sətirdən çox ola bilməz.")
    return parsed


def parse_asset_file(uploaded) -> list[dict]:
    name = (getattr(uploaded, "name", "") or "").lower()
    if name.endswith(".xlsx"):
        return parse_asset_xlsx(uploaded)
    if name.endswith(".docx"):
        return parse_asset_docx(uploaded)
    raise ValueError("Yalnız .xlsx və ya .docx fayl qəbul olunur.")


def preview_rows(parsed, floors) -> list[dict]:
    floors = list(floors)
    ready = []
    for item in parsed:
        name = (item.get("name") or "").strip()
        if not name:
            continue
        floor = match_floor(item.get("floor"), floors)
        location = (item.get("location") or "").strip()
        notes = (item.get("notes") or "").strip()
        floor_text = (item.get("floor") or "").strip()
        if floor is None and floor_text:
            if not location:
                location = floor_text
            elif floor_text.casefold() not in location.casefold():
                notes = f"{notes} · {floor_text}".strip(" ·")
        last = item.get("last_inspected_on")
        if isinstance(last, str):
            last = _parse_date(last)
        next_on = add_one_month(last) if last else None
        category = _category_code(item.get("category"))
        ready.append(
            {
                "name": name[:160],
                "code": (item.get("code") or "").strip()[:64],
                "category": category,
                "category_label": str(BuildingAssetCategory(category).label),
                "floor_id": floor.pk if floor else None,
                "floor_label": floor.name if floor else "",
                "location": location[:160],
                "quantity": item.get("quantity") or 1,
                "notes": notes,
                "last_inspected_on": last.isoformat() if last else "",
                "next_inspection_on": next_on.isoformat() if next_on else "",
            }
        )
    return ready


@transaction.atomic
def commit_asset_rows(rows) -> tuple[int, int]:
    created = 0
    updated = 0
    for row in rows:
        last = _parse_date(row.get("last_inspected_on")) if row.get("last_inspected_on") else None
        next_on = _parse_date(row.get("next_inspection_on")) if row.get("next_inspection_on") else None
        if last and not next_on:
            next_on = add_one_month(last)
        existing = BuildingAsset.objects.filter(
            name=row["name"],
            code=row.get("code") or "",
            is_active=True,
        ).first()
        floor = None
        if row.get("floor_id"):
            floor = Floor.objects.filter(pk=row["floor_id"]).first()
        if existing:
            existing.category = row.get("category") or BuildingAssetCategory.OTHER
            if floor:
                existing.floor = floor
            if row.get("location"):
                existing.location = row["location"]
            if row.get("quantity"):
                existing.quantity = row["quantity"]
            if row.get("notes"):
                existing.notes = row["notes"]
            if last:
                existing.last_inspected_on = last
            if next_on:
                existing.next_inspection_on = next_on
            existing.save()
            updated += 1
            continue
        BuildingAsset.objects.create(
            name=row["name"],
            code=row.get("code") or "",
            category=row.get("category") or BuildingAssetCategory.OTHER,
            floor=floor,
            location=row.get("location") or "",
            quantity=row.get("quantity") or 1,
            notes=row.get("notes") or "",
            last_inspected_on=last,
            next_inspection_on=next_on,
        )
        created += 1
    return created, updated


def asset_template_response() -> HttpResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Avadanliq"
    sheet.append(["Ad", "Kod", "Kateqoriya", "Mərtəbə", "Yer", "Say", "Son baxış", "Qeyd"])
    sheet.append(["Lift 1", "L-1", "Lift", "1-ci mərtəbə", "Hol", 1, "", ""])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="avadanliq-siyahisi.xlsx"'
    return response
