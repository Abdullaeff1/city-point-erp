"""Read a cost Excel into blocks, then append the rows the user keeps."""

from __future__ import annotations

import io
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook

from apps.core.services import DomainError
from apps.maintenance.costs import line_amount, refresh_cost_totals
from apps.maintenance.models import CostLineKind, CostLineSource, WorkOrderCostLine

MAX_ROWS = 500

_ALIASES = {
    "ad": "name",
    "mehsul": "name",
    "meyar": "name",
    "malin adi": "name",
    "isin adi": "name",
    "name": "name",
    "description": "name",
    "tesvir": "name",
    "miqdar": "qty",
    "qty": "qty",
    "quantity": "qty",
    "say": "qty",
    "vahid": "unit",
    "olcu vahidi": "unit",
    "unit": "unit",
    "qiymet": "price",
    "price": "price",
    "deyer": "price",
    "cem": "amount",
    "odenis": "amount",
}

_LABOR_NAME = {"isin adi"}
_MATERIAL_NAME = {"malin adi"}
_LABOR_SECTIONS = {"iscilik", "fehle"}


def preview_key(code: str) -> str:
    return f"wo_cost_preview:{code}"


def _norm(value) -> str:
    text = str(value or "").strip()
    text = text.replace("İ", "i").replace("I", "i")
    text = text.casefold()
    for src, dst in (
        ("ı", "i"),
        ("i̇", "i"),
        ("ə", "e"),
        ("ö", "o"),
        ("ü", "u"),
        ("ş", "s"),
        ("ç", "c"),
        ("ğ", "g"),
    ):
        text = text.replace(src, dst)
    return " ".join(text.split())


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value) -> str:
    if _blank(value):
        return ""
    return str(value).strip()


def _decimal(value):
    if _blank(value):
        return None
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, bool):
        raise ValueError("rəqəm deyil")
    elif isinstance(value, (int, float)):
        number = Decimal(str(round(value, 6)))
    else:
        text = str(value).strip().replace(" ", "").replace(",", ".")
        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError("rəqəm deyil") from exc
    return number


def _money(number: Decimal) -> str:
    return str(number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def sheet_template_response(kind: str) -> HttpResponse:
    if kind not in {CostLineKind.MATERIAL, CostLineKind.LABOR}:
        raise DomainError("Naməlum cədvəl.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Material" if kind == CostLineKind.MATERIAL else "Iscilik"
    sheet.append(["Ad", "Miqdar", "Vahid", "Qiymət"])
    if kind == CostLineKind.MATERIAL:
        sheet.append(["Nasos", 2, "ədəd", 15])
        filename = "material-siyahisi.xlsx"
    else:
        sheet.append(["Usta işi", 3, "saat", 20])
        filename = "iscilik-siyahisi.xlsx"
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def read_cost_preview(upload, *, default_kind: str) -> dict:
    if default_kind not in {CostLineKind.MATERIAL, CostLineKind.LABOR}:
        raise DomainError("Naməlum cədvəl.")
    name = getattr(upload, "name", "") or ""
    if not name.lower().endswith(".xlsx"):
        raise DomainError("Yalnız .xlsx faylı qəbul olunur.")
    try:
        workbook = load_workbook(upload, data_only=True, read_only=True)
    except Exception as exc:
        raise DomainError("Excel faylı oxunmadı.") from exc
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            blocks = _parse_sheet(worksheet.iter_rows(values_only=True), default_kind=default_kind)
            if blocks:
                sheets.append({"name": worksheet.title, "blocks": blocks})
    finally:
        workbook.close()
    if not sheets:
        raise DomainError("Əlavə olunacaq sətir yoxdur.")
    return {"default_kind": default_kind, "sheets": sheets}


def labeled_description(block: dict, row: dict) -> str:
    title = (block.get("title") or "").strip()
    name = row["description"]
    if not title:
        return name[:255]
    return f"{title} — {name}"[:255]


def _parse_sheet(rows, *, default_kind: str) -> list[dict]:
    parsed_rows = list(rows)
    if not any(_header(row) for row in parsed_rows):
        return _plain_sheet(parsed_rows, default_kind=default_kind)
    blocks = []
    section = ""
    current = None
    header_map = None
    count = 0
    for row_number, raw in enumerate(parsed_rows, start=1):
        found = _header(raw)
        if found:
            mapping, name_label = found
            current = {
                "title": section,
                "kind": _block_kind(name_label, section, default_kind),
                "rows": [],
            }
            blocks.append(current)
            header_map = mapping
            continue
        if current is not None and header_map is not None:
            item = _data_row(raw, header_map, row_number=row_number, kind=current["kind"])
            if item is not None:
                count += 1
                if count > MAX_ROWS:
                    raise DomainError("Ən çoxu 500 sətir yüklənə bilər.")
                current["rows"].append(item)
                continue
        title = _section_title(raw)
        if title:
            current = None
            header_map = None
            section = title
    return [block for block in blocks if block["rows"]]


def _plain_sheet(rows, *, default_kind: str) -> list[dict]:
    """A sheet with no header is Ad, Miqdar, Vahid, Qiymət from the first cell."""
    header = {"name": 0, "qty": 1, "unit": 2, "price": 3}
    block = {"title": "", "kind": default_kind, "rows": []}
    for row_number, raw in enumerate(rows, start=1):
        item = _data_row(raw, header, row_number=row_number, kind=default_kind)
        if item is None:
            continue
        if len(block["rows"]) >= MAX_ROWS:
            raise DomainError("Ən çoxu 500 sətir yüklənə bilər.")
        block["rows"].append(item)
    if not block["rows"]:
        return []
    return [block]


def _header(row):
    mapping = {}
    name_label = ""
    for index, cell in enumerate(row or ()):
        key = _ALIASES.get(_norm(cell))
        if not key or key in mapping:
            if key == "name" and not name_label:
                name_label = _norm(cell)
            continue
        mapping[key] = index
        if key == "name":
            name_label = _norm(cell)
    if "name" in mapping and ({"qty", "price", "amount"} & set(mapping)):
        return mapping, name_label
    return None


def _section_title(row) -> str:
    cells = [cell for cell in (row or ()) if not _blank(cell)]
    if len(cells) != 1:
        return ""
    text = _text(cells[0])
    norm = _norm(text)
    if not norm or norm.startswith("qeyd") or norm in _ALIASES:
        return ""
    try:
        if _decimal(text) is not None:
            return ""
    except ValueError:
        pass
    return text[:80]


def _block_kind(name_label: str, section: str, default_kind: str) -> str:
    if name_label in _LABOR_NAME:
        return CostLineKind.LABOR
    if name_label in _MATERIAL_NAME:
        return CostLineKind.MATERIAL
    if _norm(section) in _LABOR_SECTIONS:
        return CostLineKind.LABOR
    return default_kind


def _data_row(raw, header_map, *, row_number: int, kind: str):
    cells = list(raw or ())

    def cell(key):
        index = header_map.get(key)
        if index is None or index >= len(cells):
            return None
        return cells[index]

    description = _text(cell("name"))
    if not description or _norm(description).startswith("qeyd"):
        return None
    if _norm(description) in _ALIASES:
        return None
    try:
        qty = _decimal(cell("qty"))
    except ValueError as exc:
        raise DomainError(f"Sətir {row_number}: miqdar rəqəm olmalıdır.") from exc
    if qty is None:
        qty = Decimal("1")
    if qty <= 0:
        raise DomainError(f"Sətir {row_number}: miqdar sıfırdan böyük olmalıdır.")
    try:
        price = _decimal(cell("price"))
        amount = _decimal(cell("amount"))
    except ValueError as exc:
        raise DomainError(f"Sətir {row_number}: qiymət rəqəm olmalıdır.") from exc
    if price is not None and price < 0:
        raise DomainError(f"Sətir {row_number}: qiymət mənfi ola bilməz.")
    if price is None and amount is not None and amount > 0 and qty > 0:
        price = amount / qty
    if price is None:
        price = Decimal("0")
    default_unit = "saat" if kind == CostLineKind.LABOR else "ədəd"
    unit = _text(cell("unit")) or default_unit
    return {
        "description": description[:255],
        "qty": str(qty.quantize(Decimal("0.001"))),
        "unit": unit[:16],
        "unit_price": _money(price),
    }


@transaction.atomic
def append_cost_rows(wo, *, kind: str, rows: list[dict]) -> int:
    if kind not in {CostLineKind.MATERIAL, CostLineKind.LABOR}:
        raise DomainError("Naməlum cədvəl.")
    last = (
        WorkOrderCostLine.objects.filter(work_order=wo, kind=kind)
        .order_by("-sort_order")
        .values_list("sort_order", flat=True)
        .first()
    )
    sort_order = (last or 0) + 1
    for row in rows:
        qty = Decimal(row["qty"])
        unit_price = Decimal(row["unit_price"])
        WorkOrderCostLine.objects.create(
            work_order=wo,
            kind=kind,
            source=CostLineSource.MANUAL,
            description=row["description"][:255],
            qty=qty,
            unit=(row.get("unit") or "ədəd")[:16],
            unit_price=unit_price,
            amount=line_amount(qty, unit_price),
            sort_order=sort_order,
        )
        sort_order += 1
    refresh_cost_totals(wo)
    return len(rows)
