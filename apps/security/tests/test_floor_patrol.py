from datetime import date, datetime, time
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.db import IntegrityError
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from docx import Document
from django.utils.timezone import make_aware

from apps.accounts.models import Role
from apps.audit.models import AuditLog
from apps.residents.models import ResidentCompany, ResidentEmployee
from apps.security.models import FloorPatrolRound, FloorPatrolSheet, PatrolRoundStatus
from apps.security.patrol import (
    open_or_create_sheet,
    round_calendar_date,
    suggested_night_date,
)

User = get_user_model()


class FloorPatrolTests(TestCase):
    def setUp(self):
        self.security = User.objects.create_user(
            username="patrol-sec",
            email="security@patrol.az",
            password="Pass12345!",
            role=Role.SECURITY,
        )
        self.admin = User.objects.create_user(
            username="patrol-admin",
            email="admin@patrol.az",
            password="Pass12345!",
            role=Role.ADMIN,
            is_staff=True,
        )
        self.management = User.objects.create_user(
            username="patrol-mgmt",
            email="manager@patrol.az",
            password="Pass12345!",
            role=Role.MANAGEMENT,
        )
        self.reception = User.objects.create_user(
            username="patrol-rec",
            email="reception@patrol.az",
            password="Pass12345!",
            role=Role.RECEPTION,
            is_staff=True,
        )
        self.resident = User.objects.create_user(
            username="patrol-res",
            email="resident@patrol.az",
            password="Pass12345!",
            role=Role.RESIDENT_USER,
        )
        company = ResidentCompany.objects.create(
            name="City Point", slug="city-point", is_internal=True
        )
        other = ResidentCompany.objects.create(name="Başqa MMC", slug="basqa-patrol")
        self.nicat = ResidentEmployee.objects.create(company=company, full_name="Nicat")
        self.elvin = ResidentEmployee.objects.create(company=company, full_name="Elvin")
        self.miryaqub = ResidentEmployee.objects.create(company=company, full_name="Miryaqub")
        ResidentEmployee.objects.create(company=other, full_name="Kənar Əməkdaş")
        self.client = Client()

    def _open(self, floors=("2", "7", "12")):
        sheet, _created = open_or_create_sheet(
            night_date=date(2026, 10, 20),
            floor_codes=list(floors),
            user=self.security,
        )
        return sheet

    def test_midnight_round_belongs_to_evening_sheet(self):
        sheet = self._open()
        self.assertEqual(round_calendar_date(sheet.night_date, "00:00"), date(2026, 10, 21))
        self.client.force_login(self.security)
        self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {
                "action": "save",
                "round__2__0000__started": "00:14",
                "round__2__0000__ended": "00:15",
                "round__2__0000__inspector": "Miryaqub",
            },
        )
        rnd = FloorPatrolRound.objects.get(sheet=sheet, floor_code="2", slot="00:00")
        self.assertEqual(rnd.started_at, time(0, 14))
        self.assertEqual(rnd.ended_at, time(0, 15))
        self.assertEqual(rnd.status, PatrolRoundStatus.COMPLETED)
        self.assertEqual(sheet.night_date, date(2026, 10, 20))
        page = self.client.get(reverse("security:patrol_detail", args=[sheet.pk]))
        self.assertContains(page, "21.10.2026")
        self.assertContains(page, "00:14")
        self.assertNotContains(page, 'type="time"')
        self.assertContains(page, 'placeholder="20:32"')
        self.assertContains(page, ">Miryaqub</option>")
        self.assertNotContains(page, "Kənar Əməkdaş")

    def test_unknown_inspector_is_rejected(self):
        sheet = self._open(floors=["2"])
        self.client.force_login(self.security)
        response = self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {
                "action": "save",
                "round__2__2000__started": "20:32",
                "round__2__2000__ended": "20:33",
                "round__2__2000__inspector": "Kənar adam",
            },
        )
        self.assertContains(response, "City Point əməkdaşı olmalıdır")
        rnd = FloorPatrolRound.objects.get(sheet=sheet, floor_code="2", slot="20:00")
        self.assertEqual(rnd.inspector_name, "")

    def test_missed_keeps_start_and_clears_end(self):
        sheet = self._open()
        self.client.force_login(self.security)
        self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {
                "action": "save",
                "round__7__2200__started": "22:00",
                "round__7__2200__ended": "22:10",
                "round__7__2200__missed": "1",
                "round__7__2200__inspector": "Elvin",
            },
        )
        rnd = FloorPatrolRound.objects.get(sheet=sheet, floor_code="7", slot="22:00")
        self.assertEqual(rnd.status, PatrolRoundStatus.MISSED)
        self.assertEqual(rnd.started_at, time(22, 0))
        self.assertIsNone(rnd.ended_at)

    def test_start_after_end_names_floor_and_slot(self):
        sheet = self._open()
        self.client.force_login(self.security)
        response = self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {
                "action": "save",
                "round__7__2200__started": "22:30",
                "round__7__2200__ended": "22:00",
                "round__7__2200__inspector": "Elvin",
            },
        )
        self.assertContains(response, "Mərtəbə 7, tur 22:00")
        rnd = FloorPatrolRound.objects.get(sheet=sheet, floor_code="7", slot="22:00")
        self.assertEqual(rnd.status, PatrolRoundStatus.WAITING)
        self.assertIsNone(rnd.started_at)

    def test_month_grid_marks_checked_floors(self):
        sheet = self._open(floors=["2", "7"])
        open_or_create_sheet(
            night_date=date(2026, 10, 5), floor_codes=["7"], user=self.security
        )
        open_or_create_sheet(
            night_date=date(2026, 11, 2), floor_codes=["12"], user=self.security
        )
        self.assertEqual(sheet.status, "draft")
        self.assertFalse(
            FloorPatrolRound.objects.filter(sheet=sheet).exclude(status=PatrolRoundStatus.WAITING).exists()
        )
        self.client.force_login(self.security)
        page = self.client.get(reverse("security:patrol_detail", args=[sheet.pk]))
        self.assertContains(page, "YOXLANILDI")
        self.assertContains(page, "01.10.2026")
        self.assertContains(page, "31.10.2026")
        self.assertContains(
            page,
            'data-date="2026-10-20" data-floor="2">X</td>',
        )
        self.assertContains(
            page,
            'data-date="2026-10-20" data-floor="7">X</td>',
        )
        self.assertContains(
            page,
            'data-date="2026-10-05" data-floor="7">X</td>',
        )
        self.assertContains(
            page,
            'data-date="2026-10-20" data-floor="12"></td>',
        )
        self.assertNotContains(page, "02.11.2026")
        self.assertNotContains(page, 'name="mark__')

    def test_word_export_matches_floor_sheet_and_month_marks(self):
        sheet = self._open(floors=["2", "7"])
        self.client.force_login(self.security)
        self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {
                "action": "save",
                "round__2__2000__started": "20:32",
                "round__2__2000__ended": "20:33",
                "round__2__2000__inspector": str(self.nicat.pk),
                "round__7__2200__started": "22:00",
                "round__7__2200__missed": "1",
                "round__7__2200__inspector": str(self.elvin.pk),
            },
        )
        response = self.client.get(reverse("security:patrol_word", args=[sheet.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("wordprocessingml", response["Content-Type"])
        self.assertIn("yoxlama-2026-10-20.docx", response["Content-Disposition"])
        document = Document(BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("MƏRTƏBƏLƏRİN YOXLANILMASI", text)
        self.assertIn("YOXLANILDI", text)
        cells = [
            paragraph.text
            for table in document.tables
            for row in table.rows
            for paragraph in row.cells
        ]
        joined = "\n".join(cells)
        self.assertIn("Mərtəbə    2", joined)
        self.assertIn("Mərtəbə    7", joined)
        self.assertIn("20:32", joined)
        self.assertIn("Nicat", joined)
        self.assertIn("YOX", joined)
        self.assertIn("Elvin", joined)
        self.assertIn("X", joined)
        self.client.force_login(self.reception)
        denied = self.client.get(reverse("security:patrol_word", args=[sheet.pk]))
        self.assertEqual(denied.status_code, 403)

    def test_management_can_view_but_not_save(self):
        sheet = self._open()
        self.client.force_login(self.management)
        page = self.client.get(reverse("security:patrol_detail", args=[sheet.pk]))
        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, "Hamısını saxla")
        denied = self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {"action": "save", "round__2__2000__started": "20:32"},
        )
        self.assertEqual(denied.status_code, 403)

    def test_reception_and_resident_are_forbidden(self):
        self.client.force_login(self.reception)
        self.assertEqual(self.client.get(reverse("security:patrol_list")).status_code, 403)
        self.client.force_login(self.resident)
        self.assertEqual(self.client.get(reverse("security:patrol_list")).status_code, 403)

    def test_same_night_returns_existing_sheet(self):
        first, created = open_or_create_sheet(
            night_date=date(2026, 10, 20), floor_codes=["2"], user=self.security
        )
        second, again = open_or_create_sheet(
            night_date=date(2026, 10, 20), floor_codes=["12"], user=self.admin
        )
        self.assertTrue(created)
        self.assertFalse(again)
        self.assertEqual(first.pk, second.pk)
        self.assertEqual(FloorPatrolSheet.objects.count(), 1)
        self.assertEqual(
            FloorPatrolRound.objects.filter(sheet=first, floor_code="2").count(), 6
        )
        self.assertFalse(FloorPatrolRound.objects.filter(sheet=first, floor_code="12").exists())

    def test_integrity_error_returns_existing_sheet(self):
        sheet, _created = open_or_create_sheet(
            night_date=date(2026, 10, 20), floor_codes=["2"], user=self.security
        )
        with patch.object(
            FloorPatrolSheet.objects, "get_or_create", side_effect=IntegrityError
        ):
            again, created = open_or_create_sheet(
                night_date=date(2026, 10, 20), floor_codes=["7"], user=self.security
            )
        self.assertFalse(created)
        self.assertEqual(again.pk, sheet.pk)
        self.assertEqual(FloorPatrolSheet.objects.count(), 1)

    def test_closed_edit_is_audited_separately_from_inspector_name(self):
        sheet = self._open(floors=["2"])
        self.client.force_login(self.security)
        self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {"action": "close"},
        )
        self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {
                "action": "save",
                "round__2__2000__started": "20:32",
                "round__2__2000__ended": "20:33",
                "round__2__2000__inspector": "Miryaqub",
            },
        )
        logs = AuditLog.objects.filter(action="floor_patrol.updated").order_by("id")
        self.assertGreaterEqual(logs.count(), 2)
        latest = logs.last()
        self.assertEqual(latest.actor_id, self.security.pk)
        self.assertEqual(latest.source, "security")
        self.assertIn("Miryaqub", str(latest.new_values))
        self.assertNotEqual(latest.actor.email, "Miryaqub")

    def test_same_slot_is_not_duplicated(self):
        sheet = self._open(floors=["12"])
        self.client.force_login(self.security)
        self.client.post(
            reverse("security:patrol_detail", args=[sheet.pk]),
            {"action": "add_floor", "floor_code": "12"},
        )
        self.assertEqual(
            FloorPatrolRound.objects.filter(sheet=sheet, floor_code="12").count(), 6
        )

    def test_cutoff_uses_baku_clock_and_setting(self):
        self.assertEqual(
            suggested_night_date(make_aware(datetime(2026, 10, 21, 7, 59))),
            date(2026, 10, 20),
        )
        self.assertEqual(
            suggested_night_date(make_aware(datetime(2026, 10, 21, 8, 0))),
            date(2026, 10, 21),
        )

    @override_settings(FLOOR_PATROL_NIGHT_CUTOFF="09:30")
    def test_cutoff_setting_can_move(self):
        self.assertEqual(
            suggested_night_date(make_aware(datetime(2026, 10, 21, 9, 0))),
            date(2026, 10, 20),
        )
