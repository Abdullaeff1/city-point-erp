"""Word export of a night patrol sheet, laid out like the handwritten YOXLAMA form."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from apps.security.patrol import PATROL_FLOOR_CODES, month_check_grid, posted_round_rows

_TITLE = "GECƏ VAXTI MÜHAFİZƏ XİDMƏTLƏRİ TƏRƏFİNDƏN MƏRTƏBƏLƏRİN YOXLANILMASI"


def build_patrol_docx(sheet) -> bytes:
    document = Document()
    _page(document)
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(heading, _TITLE, bold=True, size=14)

    for block in posted_round_rows(sheet, None):
        document.add_paragraph()
        _floor_table(document, block)

    document.add_paragraph()
    mark_title = document.add_paragraph()
    mark_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(mark_title, "YOXLANILDI", bold=True, size=14)
    _month_table(document, month_check_grid(sheet))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _page(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)


def _floor_table(document, block):
    rounds = block["rounds"]
    table = document.add_table(rows=3 + len(rounds), cols=5)
    table.style = "Table Grid"
    table.autofit = True
    title = table.cell(0, 0).merge(table.cell(0, 4))
    saat = table.cell(1, 1).merge(table.cell(1, 2))
    table.cell(1, 0).merge(table.cell(2, 0))
    table.cell(1, 3).merge(table.cell(2, 3))
    table.cell(1, 4).merge(table.cell(2, 4))
    _cell(title, f"Mərtəbə    {block['floor']}", bold=True, size=12)
    _cell(table.cell(1, 0), "Tarix", bold=True)
    _cell(saat, "Saat", bold=True)
    _cell(table.cell(1, 3), "Yoxlayan şəxsin A.S.", bold=True)
    _cell(table.cell(1, 4), "Qeyd", bold=True)
    _cell(table.cell(2, 1), "Başladı", bold=True)
    _cell(table.cell(2, 2), "Yekunlaşdı", bold=True)
    for index, row in enumerate(rounds):
        at = index + 3
        ended = "YOX" if row["missed"] else row["ended"]
        _cell(table.cell(at, 0), row["date"].strftime("%d.%m.%Y"))
        _cell(table.cell(at, 1), row["started"])
        _cell(table.cell(at, 2), ended)
        _cell(table.cell(at, 3), row["inspector"])
        _cell(table.cell(at, 4), row["note"], align=WD_ALIGN_PARAGRAPH.LEFT)


def _month_table(document, rows):
    columns = 1 + len(PATROL_FLOOR_CODES)
    table = document.add_table(rows=2 + len(rows), cols=columns)
    table.style = "Table Grid"
    floors = table.cell(0, 1).merge(table.cell(0, columns - 1))
    _cell(table.cell(0, 0), "Tarix", bold=True, size=9)
    _cell(floors, "Mərtəbələr", bold=True, size=9)
    for index, code in enumerate(PATROL_FLOOR_CODES):
        _cell(table.cell(1, index + 1), code, bold=True, size=8)
    for row_index, row in enumerate(rows):
        at = row_index + 2
        _cell(table.cell(at, 0), row["date"].strftime("%d.%m.%Y"), size=8)
        for cell_index, cell in enumerate(row["cells"]):
            _cell(table.cell(at, cell_index + 1), "X" if cell["checked"] else "", bold=True, size=9)


def _cell(cell, text, *, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.CENTER):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    if paragraph.runs:
        run = paragraph.runs[0]
        run.text = text or ""
        for extra in paragraph.runs[1:]:
            extra.text = ""
    else:
        run = paragraph.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"


def _run(paragraph, text, *, bold=False, size=11):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    return run
