"""Qonaq ziyarətləri — Excel export (günlük / həftəlik / aylıq / tarix aralığı).

Bir vərəq: şirkətlər ad sırası ilə, hər şirkətdə nömrələnmiş qonaq cədvəli.
"""

from __future__ import annotations

from collections import defaultdict
from datetime import date
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.reception.models import mask_fin


HEADERS = (
    "№",
    "Qonaq",
    "Seriya №",
    "Tarix",
    "Host",
    "Gəliş səbəbi",
    "Giriş",
    "Çıxış",
    "Vəsiqə",
    "Status",
    "Qeyd",
)

WIDTHS = [5, 24, 14, 12, 22, 22, 8, 8, 12, 14, 28]
COLS = len(HEADERS)


def _fmt_time(value) -> str:
    if not value:
        return ""
    return timezone.localtime(value).strftime("%H:%M")


def _fmt_period(date_from: date | None, date_to: date | None) -> str:
    if not date_from:
        return ""
    if not date_to or date_from == date_to:
        return date_from.strftime("%d.%m.%Y")
    return f"{date_from.strftime('%d.%m.%Y')} – {date_to.strftime('%d.%m.%Y')}"


def visits_xlsx_filename(date_from: date, date_to: date) -> str:
    if date_from == date_to:
        return f"qonaqlar-{date_from.isoformat()}.xlsx"
    return f"qonaqlar-{date_from.isoformat()}_{date_to.isoformat()}.xlsx"


def _group_by_company(visits) -> list[tuple[str, list]]:
    buckets: dict[str, list] = defaultdict(list)
    for visit in visits:
        name = visit.company.name if visit.company_id else "—"
        buckets[name].append(visit)
    names = sorted(buckets.keys(), key=lambda n: (n == "—", n.casefold()))
    return [(name, buckets[name]) for name in names]


def _box_border(thin, medium, *, left_edge=False, right_edge=False, top_edge=False, bottom_edge=False):
    return Border(
        left=medium if left_edge else thin,
        right=medium if right_edge else thin,
        top=medium if top_edge else thin,
        bottom=medium if bottom_edge else thin,
    )


def build_visits_xlsx(
    visits,
    *,
    can_sensitive: bool = False,
    date_from: date | None = None,
    date_to: date | None = None,
) -> bytes:
    title_font = Font(name="Calibri", size=14, bold=True, color="0C2745")
    period_font = Font(name="Calibri", size=11, color="475569")
    company_font = Font(name="Calibri", size=12, bold=True, color="0C2745")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=11)
    num_font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    left = Alignment(horizontal="left", vertical="center")
    center = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="94A3B8")
    medium = Side(style="medium", color="1F497D")
    title_fill = PatternFill("solid", fgColor="E8EEF5")
    company_fill = PatternFill("solid", fgColor="D6E3F0")
    header_fill = PatternFill("solid", fgColor="1F497D")
    alt_fill = PatternFill("solid", fgColor="F8FAFC")

    groups = _group_by_company(list(visits))
    wb = Workbook()
    ws = wb.active
    ws.title = "Qonaqlar"

    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    row = 1
    # Report title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COLS)
    title = ws.cell(row=row, column=1, value="Qonaqlar hesabatı")
    title.font = title_font
    title.alignment = left
    title.fill = title_fill
    for col in range(1, COLS + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = title_fill
        cell.border = _box_border(
            thin, medium, left_edge=col == 1, right_edge=col == COLS, top_edge=True
        )
    row += 1

    # Period + total
    period = _fmt_period(date_from, date_to)
    total = sum(len(v) for _, v in groups)
    subtitle = f"{period}  ·  {total} qonaq  ·  {len(groups)} şirkət" if period else f"{total} qonaq"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COLS)
    sub = ws.cell(row=row, column=1, value=subtitle)
    sub.font = period_font
    sub.alignment = left
    for col in range(1, COLS + 1):
        ws.cell(row=row, column=col).border = _box_border(
            thin, medium, left_edge=col == 1, right_edge=col == COLS, bottom_edge=True
        )
    row += 2

    if not groups:
        ws.cell(row=row, column=1, value="Seçilmiş dövr üçün qonaq yoxdur.").font = period_font
    else:
        for g_i, (company_name, company_visits) in enumerate(groups):
            if g_i > 0:
                row += 1  # blank between companies

            # Company block heading
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=COLS)
            heading = ws.cell(
                row=row,
                column=1,
                value=f"{company_name}  ({len(company_visits)})",
            )
            heading.font = company_font
            heading.alignment = left
            for col in range(1, COLS + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = company_fill
                cell.border = _box_border(
                    thin, medium, left_edge=col == 1, right_edge=col == COLS, top_edge=True
                )
            row += 1

            # Column headers
            for col, label in enumerate(HEADERS, start=1):
                cell = ws.cell(row=row, column=col, value=label)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center
                cell.border = _box_border(thin, medium, left_edge=col == 1, right_edge=col == COLS)
            row += 1

            for n, visit in enumerate(company_visits, start=1):
                guest = visit.guest
                serial = guest.fin_code if can_sensitive else mask_fin(guest.fin_code)
                values = [
                    n,
                    guest.display_name,
                    serial or "—",
                    visit.scheduled_for.strftime("%d.%m.%Y") if visit.scheduled_for else "",
                    visit.host.full_name if visit.host_id else "—",
                    visit.visit_type.name if visit.visit_type_id else "—",
                    _fmt_time(visit.check_in_at) or "—",
                    _fmt_time(visit.check_out_at) or "—",
                    str(visit.id_document_label),
                    visit.get_status_display(),
                    (visit.notes or "").strip(),
                ]
                is_last = n == len(company_visits)
                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.font = num_font if col == 1 else cell_font
                    cell.alignment = center if col in (1, 3, 4, 7, 8, 9, 10) else left
                    if n % 2 == 0:
                        cell.fill = alt_fill
                    cell.border = _box_border(
                        thin,
                        medium,
                        left_edge=col == 1,
                        right_edge=col == COLS,
                        bottom_edge=is_last,
                    )
                row += 1

    ws.freeze_panes = "A3"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def visits_xlsx_http_response(
    visits,
    date_from: date,
    date_to: date,
    *,
    can_sensitive: bool = False,
) -> HttpResponse:
    payload = build_visits_xlsx(
        visits,
        can_sensitive=can_sensitive,
        date_from=date_from,
        date_to=date_to,
    )
    filename = visits_xlsx_filename(date_from, date_to)
    response = HttpResponse(
        payload,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
