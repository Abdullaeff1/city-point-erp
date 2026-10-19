"""Parse an uploaded Excel or Word guest list and save it as pre-registered visits."""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime
from typing import Iterable

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.reception.models import GuestListImport, VisitorType
from apps.reception.services import pre_register_visit

PREVIEW_SESSION_KEY = "guest_list_preview"
MAX_ROWS = 500

_HEADER_ALIASES = {
    "ad": "first_name",
    "soyad": "last_name",
    "first name": "first_name",
    "firstname": "first_name",
    "last name": "last_name",
    "lastname": "last_name",
    "ad soyad": "full_name",
    "adsoyad": "full_name",
    "full name": "full_name",
    "fullname": "full_name",
    "name": "full_name",
    "şirkət": "company",
    "sirket": "company",
    "company": "company",
    "tarix": "date",
    "date": "date",
    "telefon": "phone",
    "phone": "phone",
    "tel": "phone",
}


def _norm_header(value) -> str:
    text = str(value or "").strip().casefold()
    text = text.replace("ı", "i").replace("İ", "i")
    return " ".join(text.split())


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def parse_guest_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _split_full_name(full_name: str) -> tuple[str, str]:
    parts = (full_name or "").strip().split(None, 1)
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], parts[1]


def _map_headers(headers: Iterable) -> dict[int, str]:
    mapped = {}
    for index, header in enumerate(headers):
        key = _HEADER_ALIASES.get(_norm_header(header))
        if key and key not in mapped.values():
            mapped[index] = key
    return mapped


def _rows_from_mapped(table_rows: list[list]) -> list[dict]:
    if not table_rows:
        return []
    mapping = _map_headers(table_rows[0])
    if "first_name" not in mapping.values() and "full_name" not in mapping.values():
        raise ValueError("Başlıq sətri tanınmadı. Ad və ya Ad Soyad sütunu lazımdır.")
    parsed = []
    for raw in table_rows[1:]:
        if not any(_cell_text(cell) for cell in raw):
            continue
        item = {"first_name": "", "last_name": "", "company_name": "", "phone": "", "scheduled_for": None}
        for index, key in mapping.items():
            if index >= len(raw):
                continue
            value = raw[index]
            if key == "first_name":
                item["first_name"] = _cell_text(value)
            elif key == "last_name":
                item["last_name"] = _cell_text(value)
            elif key == "full_name":
                first, last = _split_full_name(_cell_text(value))
                item["first_name"] = item["first_name"] or first
                item["last_name"] = item["last_name"] or last
            elif key == "company":
                item["company_name"] = _cell_text(value)
            elif key == "phone":
                item["phone"] = _cell_text(value)
            elif key == "date":
                item["scheduled_for"] = parse_guest_date(value)
                item["date_raw"] = _cell_text(value)
        parsed.append(item)
        if len(parsed) > MAX_ROWS:
            raise ValueError(f"Siyahı {MAX_ROWS} sətirdən çox ola bilməz.")
    return parsed


def parse_xlsx(file_obj) -> list[dict]:
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _rows_from_mapped(rows)


def parse_docx(file_obj) -> list[dict]:
    from docx import Document

    document = Document(file_obj)
    if document.tables:
        table = document.tables[0]
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        return _rows_from_mapped(rows)
    parsed = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        first, last = _split_full_name(text)
        parsed.append(
            {
                "first_name": first,
                "last_name": last,
                "company_name": "",
                "phone": "",
                "scheduled_for": None,
            }
        )
        if len(parsed) > MAX_ROWS:
            raise ValueError(f"Siyahı {MAX_ROWS} sətirdən çox ola bilməz.")
    return parsed


def parse_guest_list_file(uploaded) -> list[dict]:
    name = (getattr(uploaded, "name", "") or "").lower()
    if name.endswith(".xlsx"):
        return parse_xlsx(uploaded)
    if name.endswith(".docx"):
        return parse_docx(uploaded)
    raise ValueError("Yalnız .xlsx və ya .docx fayl qəbul olunur.")


def resolve_guest_rows(rows, *, companies, company_id=None, scheduled_for=None) -> list[dict]:
    """Apply form overrides and mark rows that cannot be saved."""
    by_id = {company.pk: company for company in companies}
    by_name = {company.name.casefold().strip(): company for company in companies}
    override = by_id.get(company_id) if company_id else None
    resolved = []
    for row in rows:
        item = {
            "first_name": (row.get("first_name") or "").strip(),
            "last_name": (row.get("last_name") or "").strip(),
            "phone": (row.get("phone") or "").strip(),
            "company_id": None,
            "company_name": (row.get("company_name") or "").strip(),
            "scheduled_for": None,
            "error": "",
        }
        if override is not None:
            item["company_id"] = override.pk
            item["company_name"] = override.name
        elif item["company_name"]:
            match = by_name.get(item["company_name"].casefold())
            if match is None:
                item["error"] = "Şirkət tapılmadı."
            else:
                item["company_id"] = match.pk
                item["company_name"] = match.name
        else:
            item["error"] = "Şirkət seçilməyib."

        if scheduled_for is not None:
            item["scheduled_for"] = scheduled_for.isoformat()
        elif row.get("scheduled_for"):
            item["scheduled_for"] = row["scheduled_for"].isoformat()
        elif row.get("date_raw"):
            item["error"] = item["error"] or "Tarix tanınmadı."
        else:
            item["error"] = item["error"] or "Tarix seçilməyib."

        if not item["first_name"]:
            item["error"] = item["error"] or "Ad boşdur."
        resolved.append(item)
    return resolved


@transaction.atomic
def commit_guest_list(rows, *, filename: str, actor=None) -> list[GuestListImport]:
    """Create one import per company and date. Invalid rows are ignored."""
    valid = [row for row in rows if not row.get("error") and row.get("company_id") and row.get("scheduled_for")]
    grouped = defaultdict(list)
    for row in valid:
        grouped[(row["company_id"], row["scheduled_for"])].append(row)
    guest_type = VisitorType.objects.filter(code="guest", is_active=True).first()
    created = []
    from apps.residents.models import ResidentCompany

    companies = {
        company.pk: company
        for company in ResidentCompany.objects.filter(pk__in=[key[0] for key in grouped])
    }
    for (company_id, day_text), group_rows in grouped.items():
        company = companies.get(company_id)
        if company is None:
            continue
        day = date.fromisoformat(day_text)
        list_import = GuestListImport.objects.create(
            company=company,
            scheduled_for=day,
            filename=(filename or "")[:255],
            created_by=actor if getattr(actor, "is_authenticated", False) else None,
        )
        for row in group_rows:
            visit = pre_register_visit(
                company=company,
                first_name=row["first_name"],
                last_name=row["last_name"],
                phone=row.get("phone") or "",
                scheduled_for=day,
                visit_type=guest_type,
                portal_user=actor if getattr(actor, "is_authenticated", False) else None,
            )
            visit.list_import = list_import
            visit.save(update_fields=["list_import"])
        created.append(list_import)
    return created


def guest_list_template_response() -> HttpResponse:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Qonaqlar"
    sheet.append(["Ad", "Soyad", "Şirkət", "Tarix", "Telefon"])
    sheet.append(["Aysel", "Məmmədova", "ASBC", timezone.localdate().strftime("%d.%m.%Y"), ""])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="qonaq-siyahisi.xlsx"'
    return response
