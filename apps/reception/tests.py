from django.contrib.auth import get_user_model
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import Role
from apps.core.services import DomainError
from apps.reception.models import (
    Guest,
    GuestVisit,
    VisitStatus,
    VisitorAccess,
    VisitorAccessStatus,
    VisitorType,
    mask_fin,
    normalize_fin,
)
from apps.reception.services import ReceptionService, pre_register_visit
from apps.residents.models import ResidentCompany, ResidentEmployee
from apps.rbac.services import seed_rbac_catalog, user_has_permission


User = get_user_model()


class FinHelpersTests(TestCase):
    def test_normalize_and_mask(self):
        self.assertEqual(normalize_fin(" 7a3-b2c "), "7A3B2C")
        self.assertEqual(mask_fin("7A3B2C1D2K"), "7A3****2K")


class ReceptionServiceTests(TestCase):
    def setUp(self):
        seed_rbac_catalog()
        self.company = ResidentCompany.objects.create(name="Test Co", slug="test-co", status="active")
        self.host = ResidentEmployee.objects.create(
            company=self.company, full_name="Host Person", is_active=True
        )
        self.other = ResidentCompany.objects.create(name="Other", slug="other", status="active")
        self.other_host = ResidentEmployee.objects.create(
            company=self.other, full_name="Other Host", is_active=True
        )
        self.vtype, _ = VisitorType.objects.get_or_create(
            code="meeting", defaults={"name": "Meeting room", "sort_order": 50}
        )
        self.user = User.objects.create_user(
            username="reception_test",
            email="reception@test.az",
            password="x",
            role=Role.RECEPTION,
        )

    def test_serial_required_when_id_held(self):
        with self.assertRaises(DomainError) as ctx:
            ReceptionService.register_walk_in(
                fin_code="",
                first_name="Ali",
                last_name="Veli",
                company=self.company,
                id_document_held=True,
                guest_card_number="GC-1",
                actor=self.user,
            )
        self.assertIn("Seriya", str(ctx.exception))

    def test_walk_in_courier_without_serial_ok(self):
        """Kuryer: seriya boş, vəsiqə saxlanılmır — yalnız ad/soyad/şirkət."""
        courier, _ = VisitorType.objects.get_or_create(
            code="courier", defaults={"name": "Kuryer", "sort_order": 10}
        )
        visit = ReceptionService.register_walk_in(
            fin_code="",
            first_name="Kuryer",
            last_name="Ali",
            company=self.company,
            host=self.host,
            visit_type=courier,
            id_document_held=False,
            notes="Kuryer",
            actor=self.user,
        )
        self.assertEqual(visit.status, VisitStatus.INSIDE)
        self.assertEqual(visit.guest.fin_code, "")
        self.assertFalse(visit.id_document_held)
        self.assertEqual(visit.visit_type.code, "courier")

    def test_update_visit_adds_serial_without_changing_times(self):
        courier, _ = VisitorType.objects.get_or_create(
            code="courier", defaults={"name": "Kuryer", "sort_order": 10}
        )
        visit = ReceptionService.register_walk_in(
            fin_code="",
            first_name="Seid",
            last_name="Abdullayev",
            company=self.company,
            host=self.host,
            visit_type=courier,
            id_document_held=False,
            notes="köhnə",
            actor=self.user,
        )
        check_in = visit.check_in_at
        client = Client()
        client.login(email="reception@test.az", password="x")
        page = client.get(reverse("erp:reception"))
        self.assertContains(page, "Yadda saxla")
        res = client.post(
            reverse("erp:reception_visit_update", args=[visit.pk]),
            {
                "fin_code": "ZZ998877",
                "first_name": "Seid",
                "last_name": "Abdullayev",
                "company": self.company.pk,
                "host": self.host.pk,
                "visit_type": courier.pk,
                "notes": "yeniləndi",
            },
        )
        self.assertEqual(res.status_code, 302)
        visit.refresh_from_db()
        visit.guest.refresh_from_db()
        self.assertEqual(visit.guest.fin_code, "ZZ998877")
        self.assertEqual(visit.notes, "yeniləndi")
        self.assertEqual(visit.host_id, self.host.pk)
        self.assertEqual(visit.check_in_at, check_in)
        self.assertIsNone(visit.check_out_at)

    def test_walk_in_without_id_ok(self):
        visit = ReceptionService.register_walk_in(
            fin_code="7A3B2C1D2K",
            first_name="Ali",
            last_name="Veli",
            company=self.company,
            host=self.host,
            visit_type=self.vtype,
            id_document_held=False,
            actor=self.user,
        )
        self.assertEqual(visit.status, VisitStatus.INSIDE)
        self.assertTrue(visit.check_in_at)
        self.assertFalse(visit.id_document_held)
        self.assertFalse(VisitorAccess.objects.filter(visit=visit).exists())

    def test_walk_in_with_id_issues_guest_card(self):
        with self.assertRaises(DomainError):
            ReceptionService.register_walk_in(
                fin_code="7A3B2C1D2K",
                first_name="Ali",
                last_name="Veli",
                company=self.company,
                host=self.host,
                visit_type=self.vtype,
                id_document_held=True,
                actor=self.user,
            )
        visit = ReceptionService.register_walk_in(
            fin_code="7A3B2C1D2K",
            first_name="Ali",
            last_name="Veli",
            company=self.company,
            host=self.host,
            visit_type=self.vtype,
            id_document_held=True,
            guest_card_number="GK-1001",
            actor=self.user,
        )
        self.assertEqual(visit.status, VisitStatus.INSIDE)
        self.assertTrue(visit.id_document_held)
        self.assertTrue(hasattr(visit, "visitor_access"))
        self.assertEqual(visit.visitor_access.status, VisitorAccessStatus.ACTIVE)
        self.assertEqual(visit.visitor_access.credential, "GK-1001")

    def test_fin_reuses_guest(self):
        g1 = ReceptionService.find_or_create_guest(
            fin_code="ABC12345XY", first_name="A", last_name="B", actor=self.user
        )
        g2 = ReceptionService.find_or_create_guest(
            fin_code="abc12345xy", first_name="A", last_name="B", actor=self.user
        )
        self.assertEqual(g1.pk, g2.pk)

    def test_forgotten_card_clears_host(self):
        from apps.erp.forms import GuestVisitForm

        forgotten, _ = VisitorType.objects.get_or_create(
            code="forgotten-card",
            defaults={"name": "Kartın unutmuş əməkdaş", "sort_order": 20},
        )
        form = GuestVisitForm(
            data={
                "first_name": "A",
                "last_name": "B",
                "company": self.company.pk,
                "host": self.host.pk,
                "visit_type": forgotten.pk,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertIsNone(form.cleaned_data["host"])
        rendered = str(form["visit_type"])
        self.assertIn('data-no-host="1"', rendered)

    def test_lease_area_sets_city_point_and_clears_host(self):
        from apps.erp.forms import GuestVisitForm

        city = ResidentCompany.objects.create(
            name="City Point", slug="city-point", status="active"
        )
        lease, _ = VisitorType.objects.get_or_create(
            code="lease-area",
            defaults={"name": "İcarə sahə üçün", "sort_order": 40},
        )
        form = GuestVisitForm(
            data={
                "first_name": "A",
                "last_name": "B",
                "company": self.company.pk,
                "host": self.host.pk,
                "visit_type": lease.pk,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["company"], city)
        self.assertIsNone(form.cleaned_data["host"])
        self.assertIn('data-lock-company="city-point"', str(form["visit_type"]))

    def test_meeting_room_clears_host(self):
        from apps.erp.forms import GuestVisitForm

        form = GuestVisitForm(
            data={
                "first_name": "A",
                "last_name": "B",
                "company": self.company.pk,
                "host": self.host.pk,
                "visit_type": self.vtype.pk,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertIsNone(form.cleaned_data["host"])
        self.assertEqual(form.cleaned_data["company"], self.company)
        rendered = str(form["visit_type"])
        marker = f'value="{self.vtype.pk}"'
        option = next(part for part in rendered.split("<option") if marker in part)
        self.assertIn('data-no-host="1"', option)
        self.assertNotIn("data-lock-company", option)

    def test_technician_clears_host(self):
        from apps.erp.forms import GuestVisitForm

        technician, _ = VisitorType.objects.get_or_create(
            code="technician",
            defaults={"name": "Usta/texniki işçi", "sort_order": 70},
        )
        form = GuestVisitForm(
            data={
                "first_name": "A",
                "last_name": "B",
                "company": self.company.pk,
                "host": self.host.pk,
                "visit_type": technician.pk,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertIsNone(form.cleaned_data["host"])
        self.assertEqual(form.cleaned_data["company"], self.company)
        rendered = str(form["visit_type"])
        marker = f'value="{technician.pk}"'
        option = next(part for part in rendered.split("<option") if marker in part)
        self.assertIn("Usta/texniki işçi", option)
        self.assertIn('data-no-host="1"', option)
        self.assertNotIn("data-lock-company", option)

    def test_new_employee_keeps_host(self):
        from apps.erp.forms import GuestVisitForm

        new_employee, _ = VisitorType.objects.get_or_create(
            code="new-employee",
            defaults={"name": "Əməkdaş", "sort_order": 90},
        )
        if new_employee.name != "Əməkdaş":
            new_employee.name = "Əməkdaş"
            new_employee.save(update_fields=["name"])
        form = GuestVisitForm(
            data={
                "first_name": "A",
                "last_name": "B",
                "company": self.company.pk,
                "host": self.host.pk,
                "visit_type": new_employee.pk,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["host"], self.host)
        self.assertEqual(form.cleaned_data["visit_type"], new_employee)
        rendered = str(form["visit_type"])
        marker = f'value="{new_employee.pk}"'
        option = next(part for part in rendered.split("<option") if marker in part)
        self.assertIn("Əməkdaş", option)
        self.assertNotIn("Yeni Əməkdaş", option)
        self.assertNotIn("data-no-host", option)

    def test_kassa_and_document_reasons_are_choices(self):
        from apps.erp.forms import GuestVisitForm

        rendered = str(GuestVisitForm()["visit_type"])
        self.assertIn("Kassa", rendered)
        self.assertIn("Sənəd ilə bağlı", rendered)
        self.assertNotIn("Yeni Əməkdaş", rendered)

    def test_interview_keeps_host(self):
        from apps.erp.forms import GuestVisitForm

        interview, _ = VisitorType.objects.get_or_create(
            code="interview",
            defaults={"name": "Müsahibə", "sort_order": 100, "is_active": True},
        )
        if not interview.is_active or interview.name != "Müsahibə":
            interview.name = "Müsahibə"
            interview.is_active = True
            interview.save(update_fields=["name", "is_active"])
        form = GuestVisitForm(
            data={
                "first_name": "A",
                "last_name": "B",
                "company": self.company.pk,
                "host": self.host.pk,
                "visit_type": interview.pk,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["host"], self.host)
        self.assertEqual(form.cleaned_data["visit_type"], interview)
        rendered = str(form["visit_type"])
        marker = f'value="{interview.pk}"'
        option = next(part for part in rendered.split("<option") if marker in part)
        self.assertIn("Müsahibə", option)
        self.assertNotIn("data-no-host", option)

    def test_psychologist_company_sets_patient(self):
        from apps.erp.forms import GuestVisitForm

        psy = ResidentCompany.objects.create(
            name="PSİXOLOQ", slug="psixoloq-ax40", status="active"
        )
        patient, _ = VisitorType.objects.get_or_create(
            code="patient",
            defaults={"name": "Pasiyent", "sort_order": 80},
        )
        leyla = ResidentEmployee.objects.create(
            company=psy, full_name="LEYLA TƏBRİZ ŞAMİLOVA", is_active=True
        )
        form = GuestVisitForm(
            data={
                "first_name": "A",
                "last_name": "B",
                "company": psy.pk,
                "host": self.host.pk,
                "visit_type": self.vtype.pk,
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["company"], psy)
        self.assertEqual(form.cleaned_data["visit_type"], patient)
        self.assertEqual(form.cleaned_data["host"], leyla)
        rendered = str(form["visit_type"])
        marker = f'value="{patient.pk}"'
        option = next(part for part in rendered.split("<option") if marker in part)
        self.assertIn('data-code="patient"', option)

    def test_host_must_belong_to_company(self):
        with self.assertRaises(DomainError):
            ReceptionService.register_walk_in(
                fin_code="ZZZ11111AA",
                first_name="X",
                last_name="Y",
                company=self.company,
                host=self.other_host,
                id_document_held=True,
                guest_card_number="GK-1",
                actor=self.user,
            )

    def test_duplicate_inside_blocked(self):
        ReceptionService.register_walk_in(
            fin_code="DUP11111AA",
            first_name="Dup",
            last_name="Guest",
            company=self.company,
            id_document_held=True,
            guest_card_number="GK-DUP",
            actor=self.user,
        )
        with self.assertRaises(DomainError):
            ReceptionService.register_walk_in(
                fin_code="DUP11111AA",
                first_name="Dup",
                last_name="Guest",
                company=self.company,
                id_document_held=True,
                guest_card_number="GK-DUP2",
                actor=self.user,
            )

    def test_prereg_checkin_checkout_cancel(self):
        visit = pre_register_visit(
            company=self.company,
            first_name="Pre",
            last_name="Reg",
            host=self.host,
            visit_type=self.vtype,
        )
        self.assertEqual(visit.status, VisitStatus.PRE_REGISTERED)
        self.assertTrue(visit.invite_code)

        ReceptionService.check_in_visit(
            visit,
            actor=self.user,
            fin_code="PRE11111AA",
            id_document_held=True,
            guest_card_number="GK-PRE",
        )
        visit.refresh_from_db()
        self.assertEqual(visit.status, VisitStatus.INSIDE)

        ReceptionService.check_out_visit(visit, actor=self.user, id_returned=True)
        visit.refresh_from_db()
        self.assertEqual(visit.status, VisitStatus.LEFT)
        self.assertTrue(visit.check_out_at)

        visit2 = pre_register_visit(company=self.company, first_name="Cancel", last_name="Me")
        ReceptionService.cancel_visit(visit2, actor=self.user)
        visit2.refresh_from_db()
        self.assertEqual(visit2.status, VisitStatus.CANCELLED)
        with self.assertRaises(DomainError):
            ReceptionService.check_in_visit(
                visit2, actor=self.user, id_document_held=True, guest_card_number="GK-X"
            )

    def test_checkout_blocks_without_id_return(self):
        visit = ReceptionService.register_walk_in(
            fin_code="IDHOLD01AA",
            first_name="Id",
            last_name="Hold",
            company=self.company,
            id_document_held=True,
            guest_card_number="GK-HOLD",
            actor=self.user,
        )
        with self.assertRaises(DomainError):
            ReceptionService.check_out_visit(visit, actor=self.user, id_returned=False)
        pending = ReceptionService.check_out_visit(
            visit, actor=self.user, id_returned=False, allow_return_pending=True
        )
        self.assertEqual(pending.status, VisitStatus.RETURN_PENDING)

    def test_bulk_checkout_only_overdue_inside(self):
        from datetime import timedelta

        overdue = ReceptionService.register_walk_in(
            fin_code="",
            first_name="Old",
            last_name="Guest",
            company=self.company,
            actor=self.user,
        )
        overdue.check_in_at = timezone.now() - timedelta(hours=25)
        overdue.save(update_fields=["check_in_at"])
        held = ReceptionService.register_walk_in(
            fin_code="AA3304624",
            first_name="Seid",
            last_name="Abdullayev",
            company=self.company,
            id_document_held=True,
            guest_card_number="GK-OLD",
            actor=self.user,
        )
        held.check_in_at = timezone.now() - timedelta(hours=30)
        held.save(update_fields=["check_in_at"])
        fresh = ReceptionService.register_walk_in(
            fin_code="",
            first_name="New",
            last_name="Guest",
            company=self.company,
            actor=self.user,
        )
        client = Client()
        client.login(email="reception@test.az", password="x")
        page = client.get(reverse("erp:reception") + "?status=inside_overdue")
        self.assertContains(page, "Hamısını çıxış et")
        today = client.get(reverse("erp:reception"))
        self.assertNotContains(today, "Hamısını çıxış et")
        res = client.post(reverse("erp:reception_overdue_checkout"))
        self.assertEqual(res.status_code, 302)
        overdue.refresh_from_db()
        held.refresh_from_db()
        fresh.refresh_from_db()
        self.assertEqual(overdue.status, VisitStatus.LEFT)
        self.assertEqual(held.status, VisitStatus.LEFT)
        self.assertFalse(held.id_document_held)
        self.assertIsNotNone(held.id_document_returned_at)
        self.assertEqual(fresh.status, VisitStatus.INSIDE)


class ReceptionPermissionTests(TestCase):
    def setUp(self):
        seed_rbac_catalog()
        self.reception = User.objects.create_user(
            username="rec_perm", email="r@test.az", password="x", role=Role.RECEPTION
        )
        self.fm = User.objects.create_user(
            username="fm_perm", email="fm@test.az", password="x", role=Role.PROPERTY_FM
        )

    def test_reception_has_granular_perms(self):
        self.assertTrue(user_has_permission(self.reception, "reception.create_visit"))
        self.assertTrue(user_has_permission(self.reception, "reception.view_sensitive_data"))
        self.assertTrue(user_has_permission(self.fm, "reception.view"))
        self.assertFalse(user_has_permission(self.fm, "reception.create_visit"))


class ReceptionAPITests(TestCase):
    def setUp(self):
        seed_rbac_catalog()
        self.client = Client()
        self.user = User.objects.create_user(
            username="api_rec",
            email="api-rec@test.az",
            password="pass12345",
            role=Role.RECEPTION,
        )
        self.company = ResidentCompany.objects.create(name="API Co", slug="api-co", status="active")
        self.client.login(email="api-rec@test.az", password="pass12345")

    def test_today_endpoint(self):
        res = self.client.get("/api/reception/today/")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("stats", data)
        self.assertIn("results", data)


class ReceptionExcelExportTests(TestCase):
    def setUp(self):
        seed_rbac_catalog()
        self.client = Client()
        self.user = User.objects.create_user(
            username="rec_xlsx",
            email="rec-xlsx@test.az",
            password="pass12345",
            role=Role.RECEPTION,
        )
        self.company = ResidentCompany.objects.create(
            name="Export Co", slug="export-co", status="active"
        )
        self.vtype, _ = VisitorType.objects.get_or_create(
            code="courier", defaults={"name": "Kuryer", "sort_order": 10}
        )
        self.client.login(email="rec-xlsx@test.az", password="pass12345")
        today = timezone.localdate()
        ReceptionService.register_walk_in(
            fin_code="",
            first_name="Kuryer",
            last_name="Test",
            company=self.company,
            visit_type=self.vtype,
            id_document_held=False,
            notes="Excel test",
            actor=self.user,
        )
        self.today = today

    def test_export_xlsx_today(self):
        from io import BytesIO

        from openpyxl import load_workbook

        res = self.client.get(reverse("erp:reception_export") + f"?range=today&date={self.today.isoformat()}")
        self.assertEqual(res.status_code, 200)
        self.assertIn(
            "spreadsheetml.sheet",
            res["Content-Type"],
        )
        self.assertIn(".xlsx", res["Content-Disposition"])
        wb = load_workbook(BytesIO(res.content))
        self.assertEqual(wb.sheetnames, ["Qonaqlar"])
        ws = wb.active
        self.assertEqual(ws["A1"].value, "Qonaqlar hesabatı")
        # row1 title, row2 period, row3 blank, row4 company
        self.assertTrue(str(ws["A4"].value).startswith("Export Co"))
        self.assertEqual(ws["A5"].value, "№")
        self.assertEqual(ws["B5"].value, "Qonaq")
        self.assertEqual(ws["F5"].value, "Gəliş səbəbi")
        self.assertEqual(ws["A6"].value, 1)
        self.assertEqual(ws["B6"].value, "Kuryer Test")
        self.assertEqual(ws["F6"].value, "Kuryer")

    def test_export_xlsx_companies_alpha_one_sheet(self):
        from io import BytesIO

        from openpyxl import load_workbook

        other = ResidentCompany.objects.create(name="Beta Co", slug="beta-co", status="active")
        ReceptionService.register_walk_in(
            fin_code="",
            first_name="Qonaq",
            last_name="Beta",
            company=other,
            visit_type=self.vtype,
            id_document_held=False,
            actor=self.user,
        )
        res = self.client.get(reverse("erp:reception_export") + f"?range=today&date={self.today.isoformat()}")
        wb = load_workbook(BytesIO(res.content))
        self.assertEqual(wb.sheetnames, ["Qonaqlar"])
        ws = wb.active
        # Alpha order: Beta Co before Export Co (row4 after title/period/blank)
        self.assertTrue(str(ws["A4"].value).startswith("Beta Co"))
        self.assertEqual(ws["B6"].value, "Qonaq Beta")
        found_export = False
        for r in range(1, ws.max_row + 1):
            val = ws.cell(r, 1).value
            if isinstance(val, str) and val.startswith("Export Co"):
                found_export = True
                self.assertEqual(ws.cell(r + 2, 2).value, "Kuryer Test")
                self.assertEqual(ws.cell(r + 2, 1).value, 1)
                break
        self.assertTrue(found_export)

    def test_export_xlsx_week_range(self):
        res = self.client.get(reverse("erp:reception_export") + "?range=week")
        self.assertEqual(res.status_code, 200)
        self.assertIn(".xlsx", res["Content-Disposition"])

    def test_export_page_shows_excel_controls(self):
        res = self.client.get(reverse("erp:reception") + "?range=month")
        self.assertEqual(res.status_code, 200)
        body = res.content.decode("utf-8", "replace")
        self.assertIn("Excel yüklə", body)
        self.assertIn('data-guest-detail=', body)
        self.assertIn("guest-detail-reason", body)
        self.assertRegex(body, r'guest-detail-reason">\s*Kuryer')
        self.assertIn("Kuryer Test", body)
        self.assertNotIn("E-poçt", body)
        self.assertNotIn("Mərtəbə", body)
        self.assertIn("Bu həftə", body)
        self.assertIn("Bu ay", body)
        self.assertIn("Tarix aralığı", body)
        panel = body.split('id="journal-panel"', 1)[1]
        self.assertIn("Günlük", panel)
        self.assertIn("Bu həftə", panel)
        self.assertIn("range=week", panel)

    def test_journal_period_filters_visits(self):
        from datetime import timedelta

        past = self.today - timedelta(days=10)
        ReceptionService.register_walk_in(
            fin_code="",
            first_name="Köhnə",
            last_name="Qonaq",
            company=self.company,
            visit_type=self.vtype,
            id_document_held=False,
            actor=self.user,
        )
        from apps.reception.models import GuestVisit

        GuestVisit.objects.filter(guest__last_name="Qonaq").update(scheduled_for=past)
        today_page = self.client.get(reverse("erp:reception")).content.decode()
        self.assertIn("Kuryer Test", today_page)
        self.assertNotIn("Köhnə Qonaq", today_page)
        ranged = self.client.get(
            reverse("erp:reception") + f"?range=custom&from={past.isoformat()}&to={past.isoformat()}"
        ).content.decode()
        self.assertIn("Köhnə Qonaq", ranged)
        self.assertNotIn("Kuryer Test", ranged.split('id="journal-panel"', 1)[1])


class GuestListImportTests(TestCase):
    def setUp(self):
        seed_rbac_catalog()
        self.company = ResidentCompany.objects.create(name="Test Co", slug="list-co", status="active")
        self.user = User.objects.create_user(
            username="list_rec",
            email="list-rec@test.az",
            password="pass12345",
            role=Role.RECEPTION,
        )
        self.client = Client()
        self.client.login(email="list-rec@test.az", password="pass12345")

    def _xlsx(self, rows):
        from io import BytesIO

        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        book.save(buffer)
        buffer.seek(0)
        buffer.name = "qonaqlar.xlsx"
        return buffer

    def _docx_table(self, rows):
        from io import BytesIO

        from docx import Document

        document = Document()
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.rows[r].cells[c].text = value
        buffer = BytesIO()
        document.save(buffer)
        buffer.seek(0)
        buffer.name = "qonaqlar.docx"
        return buffer

    def test_excel_saves_preregistered_and_bulk_check_in_skips_id(self):
        from apps.reception.guest_list_import import parse_guest_list_file, resolve_guest_rows
        from apps.reception.models import GuestListImport

        uploaded = self._xlsx(
            [
                ["Ad", "Soyad", "Şirkət", "Tarix"],
                ["Aysel", "Məmmədova", "Test Co", "2026-01-02"],
                ["Kamran", "Əliyev", "Nope Ltd", "2026-01-02"],
            ]
        )
        parsed = parse_guest_list_file(uploaded)
        rows = resolve_guest_rows(
            parsed,
            companies=ResidentCompany.objects.filter(status="active"),
            scheduled_for=timezone.localdate(),
        )
        self.assertEqual(rows[0]["scheduled_for"], timezone.localdate().isoformat())
        self.assertEqual(rows[1]["error"], "Şirkət seçilməyib.")
        self.assertEqual(rows[1]["note"], "Nope Ltd")
        self.assertEqual(rows[1]["company_name"], "")

        preview = self.client.post(
            reverse("erp:reception_guest_list_preview"),
            {"guest_list": self._xlsx(
                [
                    ["Ad", "Soyad", "Şirkət", "Tarix"],
                    ["Aysel", "Məmmədova", "Test Co", "02.01.2026"],
                    ["Kamran", "Əliyev", "Nope Ltd", "02.01.2026"],
                ]
            ), "list_date": timezone.localdate().isoformat()},
        )
        self.assertEqual(preview.status_code, 302)
        page = self.client.get(preview.url)
        body = page.content.decode()
        self.assertIn("Aysel", body)
        self.assertIn("Şirkət seçilməyib.", body)
        self.assertIn("Nope Ltd", body)
        self.assertIn("Jurnala yaz", body)

        saved = self.client.post(reverse("erp:reception_guest_list_confirm"), {"action": "save"})
        self.assertEqual(saved.status_code, 302)
        journal = self.client.get(saved.url)
        self.assertIn("Hamısını giriş et", journal.content.decode())
        self.assertNotIn("Vəsiqə saxlanıldı (kart verildi)", journal.content.decode().split("Hamısını giriş et", 1)[-1])
        visits = GuestVisit.objects.filter(list_import__isnull=False)
        self.assertEqual(visits.count(), 1)
        visit = visits.get()
        self.assertEqual(visit.status, VisitStatus.PRE_REGISTERED)
        self.assertEqual(visit.guest.first_name, "Aysel")
        self.assertIsNone(visit.host_id)
        self.assertFalse(visit.id_document_held)
        self.assertEqual(visit.scheduled_for, timezone.localdate())

        list_import = GuestListImport.objects.get()
        checked = self.client.post(reverse("erp:reception_guest_list_check_in", args=[list_import.pk]))
        self.assertEqual(checked.status_code, 302)
        visit.refresh_from_db()
        self.assertEqual(visit.status, VisitStatus.INSIDE)
        self.assertFalse(visit.id_document_held)
        self.assertFalse(VisitorAccess.objects.filter(visit=visit).exists())
        self.assertIsNotNone(visit.check_in_at)

    def test_word_table_parses_names(self):
        from apps.reception.guest_list_import import parse_guest_list_file

        uploaded = self._docx_table(
            [
                ["Ad Soyad", "Şirkət", "Tarix"],
                ["Nigar Həsənova", "Test Co", "19.10.2026"],
            ]
        )
        parsed = parse_guest_list_file(uploaded)
        self.assertEqual(parsed[0]["first_name"], "Nigar")
        self.assertEqual(parsed[0]["last_name"], "Həsənova")
        self.assertEqual(parsed[0]["company_name"], "Test Co")
        self.assertEqual(parsed[0]["scheduled_for"].isoformat(), "2026-10-19")

    def test_release_sheet_reads_combined_name_and_shared_date(self):
        from datetime import datetime
        from io import BytesIO

        from openpyxl import Workbook

        from apps.reception.guest_list_import import parse_guest_list_file

        book = Workbook()
        sheet = book.active
        sheet.append(["Say", "Adı/Soyad, Atasının adı", "Təmsil etdiyi təşkilatın adı", "Gün/Saat"])
        sheet.append([None, None, None, None])
        sheet.append([1, "Kamran Məcidov", "Xalq Kapital QSC", datetime(2026, 10, 23)])
        sheet.append([2, "Bağırov Orxan", "Invest Az", None])
        sheet.merge_cells("B1:B2")
        sheet.merge_cells("D3:D4")
        buffer = BytesIO()
        book.save(buffer)
        buffer.seek(0)
        buffer.name = "buraxilis.xlsx"
        parsed = parse_guest_list_file(buffer)
        self.assertEqual(
            [(row["first_name"], row["last_name"], row["company_name"]) for row in parsed],
            [
                ("Kamran", "Məcidov", "Xalq Kapital QSC"),
                ("Bağırov", "Orxan", "Invest Az"),
            ],
        )
        self.assertTrue(all(row["scheduled_for"].isoformat() == "2026-10-23" for row in parsed))

    def test_outside_company_is_written_in_the_note(self):
        from datetime import date

        from apps.reception.guest_list_import import commit_guest_list, resolve_guest_rows
        from apps.reception.models import GuestVisit

        rows = resolve_guest_rows(
            [
                {
                    "first_name": "Kamran",
                    "last_name": "Məcidov",
                    "company_name": "Xalq Kapital QSC",
                    "scheduled_for": date(2026, 10, 23),
                }
            ],
            companies=ResidentCompany.objects.filter(status="active"),
            company_id=self.company.pk,
        )
        self.assertEqual(rows[0]["company_name"], "Test Co")
        self.assertEqual(rows[0]["note"], "Xalq Kapital QSC")
        self.assertEqual(rows[0]["error"], "")
        commit_guest_list(rows, filename="buraxilis.xlsx", actor=self.user)
        visit = GuestVisit.objects.get()
        self.assertEqual(visit.company_id, self.company.pk)
        self.assertEqual(visit.notes, "Xalq Kapital QSC")
