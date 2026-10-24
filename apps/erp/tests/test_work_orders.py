from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.accounts.models import Role
from apps.billing.models import Charge, ChargeSource
from apps.maintenance.models import CostLineKind, CostLineSource, WorkOrder, WorkOrderCostLine
from apps.procurement.models import PurchaseRequest
from apps.procurement.services import ProcurementService
from apps.property.models import Contractor
from apps.residents.models import ResidentCompany
from apps.tickets.models import TicketCategory, TicketStatus, TicketSubcategory, TicketType
from apps.tickets.seed_taxonomy import seed_ticket_taxonomy
from apps.tickets.services import create_ticket_from_portal
from apps.warehouse.models import SKU, Stock, StockMovementType, Warehouse, WarehouseIssue
from apps.warehouse.services import apply_movement

User = get_user_model()


class WorkOrderRestoreTests(TestCase):
    def setUp(self):
        seed_ticket_taxonomy()
        self.company = ResidentCompany.objects.create(
            name="Ryumka", slug="ryumka-wo", portal_active=True, is_internal=False
        )
        self.desk = User.objects.create_user(
            username="desk-wo",
            email="desk-wo@test.az",
            password="Pass12345!",
            role=Role.SERVICE_DESK,
            ticket_department_id=self._dept("servicedesk").pk,
        )
        self.reception = User.objects.create_user(
            username="reception-wo",
            email="reception-wo@test.az",
            password="Pass12345!",
            role=Role.RECEPTION,
        )
        self.admin = User.objects.create_user(
            username="admin-wo",
            email="admin-wo@test.az",
            password="Pass12345!",
            role=Role.ADMIN,
        )
        self.contractor = Contractor.objects.create(name="Podrat MMC")
        self.ticket = self._ticket("cleaning", "office-cleaning")

    def _dept(self, code):
        from apps.tickets.models import TicketDepartment

        return TicketDepartment.objects.get(code=code)

    def _ticket(self, category_slug, subcategory_slug):
        return create_ticket_from_portal(
            company=self.company,
            requester=None,
            ticket_type=TicketType.SERVICE_REQUEST,
            category=TicketCategory.objects.get(slug=category_slug),
            subcategory=TicketSubcategory.objects.get(slug=subcategory_slug),
            description="Ofis təmizliyi lazımdır",
        )

    def test_service_desk_creates_work_order_without_moving_the_ticket(self):
        self.client.force_login(self.desk)
        response = self.client.post(
            reverse("erp:ticket_detail", args=[self.ticket.code]),
            {"action": "create_wo", "contractor": self.contractor.pk},
        )
        wo = WorkOrder.objects.get(ticket=self.ticket)
        self.assertRedirects(response, reverse("erp:work_order_detail", args=[wo.code]))
        self.assertEqual(wo.contractor, self.contractor)
        self.ticket.refresh_from_db()
        self.assertEqual(self.ticket.department.code, "servicedesk")
        self.assertEqual(self.ticket.status, TicketStatus.SENT)
        page = self.client.get(reverse("erp:work_orders"))
        self.assertContains(page, wo.code)
        self.assertContains(page, "Podrat MMC")

    def test_reception_cannot_open_work_orders(self):
        self.client.force_login(self.reception)
        response = self.client.get(reverse("erp:work_orders"))
        self.assertEqual(response.status_code, 403)

    def test_stock_purchase_and_charge_stay_on_the_work_order(self):
        self.client.force_login(self.desk)
        self.client.post(
            reverse("erp:ticket_detail", args=[self.ticket.code]),
            {"action": "create_wo", "contractor": self.contractor.pk},
        )
        wo = WorkOrder.objects.get(ticket=self.ticket)
        sku = SKU.objects.create(code="NASOS", name="Nasos")
        warehouse = Warehouse.objects.create(code="MAIN", name="Əsas anbar")
        apply_movement(
            sku=sku,
            warehouse=warehouse,
            movement_type=StockMovementType.RECEIPT,
            quantity=Decimal("5"),
        )
        detail = reverse("erp:work_order_detail", args=[wo.code])
        self.client.post(
            detail,
            {"action": "issue", "sku": sku.pk, "warehouse": warehouse.pk, "quantity": "2"},
        )
        self.assertFalse(wo.stock_movements.filter(movement_type=StockMovementType.ISSUE).exists())
        self.assertEqual(Stock.objects.get(sku=sku, warehouse=warehouse).quantity, Decimal("5"))
        issue = WarehouseIssue.objects.get(work_order=wo, status="open")
        self.client.force_login(self.desk)
        denied = self.client.post(
            reverse("erp:warehouse"),
            {"action": "confirm_issue", "issue_id": issue.pk, f"issued_qty_{issue.lines.get().pk}": "2"},
        )
        self.assertEqual(denied.status_code, 403)

        line = issue.lines.get()
        self.client.force_login(self.admin)
        self.client.post(
            reverse("erp:warehouse"),
            {
                "action": "confirm_issue",
                "issue_id": issue.pk,
                f"issued_qty_{line.pk}": "2",
                f"unit_cost_{line.pk}": "4",
            },
        )
        self.assertEqual(Stock.objects.get(sku=sku, warehouse=warehouse).quantity, Decimal("3"))
        stock_line = WorkOrderCostLine.objects.get(work_order=wo, source=CostLineSource.WAREHOUSE)
        self.assertEqual(stock_line.amount, Decimal("8.00"))

        self.client.force_login(self.desk)
        self.client.post(detail, {"action": "purchase", "description": "Ehtiyat nasos", "qty": "2"})
        pr = PurchaseRequest.objects.get(work_order=wo)
        ProcurementService.approve_pr(pr, actor=self.admin)
        po = ProcurementService.create_po_from_pr(pr)
        po_line = po.lines.get()
        self.client.force_login(self.admin)
        self.client.post(
            reverse("erp:procurement"),
            {"action": "set_line_price", "line_id": po_line.pk, "unit_price": "6"},
        )
        bought = WorkOrderCostLine.objects.get(work_order=wo, source=CostLineSource.PROCUREMENT)
        self.assertEqual(bought.amount, Decimal("12.00"))

        self.client.force_login(self.desk)
        page = self.client.get(detail)
        self.assertContains(page, "wo-page")
        self.assertContains(page, "NASOS")
        self.assertContains(page, pr.code)
        self.assertContains(page, "6")
        self.client.post(
            detail,
            {
                "action": "sheet",
                "material_id": [str(stock_line.pk), str(bought.pk), ""],
                "material_description": [stock_line.description, bought.description, ""],
                "material_qty": [str(stock_line.qty), str(bought.qty), ""],
                "material_unit": [stock_line.unit, bought.unit, ""],
                "material_price": ["4", "6", ""],
                "labor_id": [""],
                "labor_description": ["Usta"],
                "labor_qty": ["1"],
                "labor_unit": ["saat"],
                "labor_price": ["10000"],
                "resident_amount": "1500",
                "billable": "1",
            },
        )
        wo.refresh_from_db()
        self.assertEqual(wo.labor_cost, Decimal("10000.00"))
        self.assertEqual(wo.resident_amount, Decimal("1500.00"))
        self.assertEqual(wo.material_cost + wo.labor_cost, Decimal("10020.00"))
        self.client.post(detail, {"action": "status", "status": "completed"})
        charge = Charge.objects.get(work_order=wo)
        self.assertEqual(charge.source, ChargeSource.WORK_ORDER)
        self.assertEqual(charge.amount, Decimal("1500.00"))

    def _xlsx(self, rows, name="siyahi.xlsx"):
        book = Workbook()
        sheet = book.active
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        book.save(buffer)
        return SimpleUploadedFile(
            name,
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_excel_appends_material_and_labor_without_replacing_rows(self):
        self.client.force_login(self.desk)
        self.client.post(
            reverse("erp:ticket_detail", args=[self.ticket.code]),
            {"action": "create_wo"},
        )
        wo = WorkOrder.objects.get(ticket=self.ticket)
        WorkOrderCostLine.objects.create(
            work_order=wo,
            kind=CostLineKind.MATERIAL,
            source=CostLineSource.MANUAL,
            description="Köhnə boya",
            qty=Decimal("1"),
            unit="ədəd",
            unit_price=Decimal("3"),
            amount=Decimal("3"),
        )
        locked = WorkOrderCostLine.objects.create(
            work_order=wo,
            kind=CostLineKind.MATERIAL,
            source=CostLineSource.WAREHOUSE,
            description="NASOS — Nasos",
            qty=Decimal("2"),
            unit="ədəd",
            unit_price=Decimal("4"),
            amount=Decimal("8"),
        )
        detail = reverse("erp:work_order_detail", args=[wo.code])
        material = self._xlsx(
            [
                ["Məhsul", "Miqdar", "Vahid", "Qiymət"],
                ["Boya", 2, "ədəd", "15,50"],
                ["Fırça", "", "", ""],
            ]
        )
        preview_url = reverse("erp:work_order_import", args=[wo.code])
        response = self.client.post(detail, {"action": "import_material", "material_file": material})
        self.assertRedirects(response, preview_url)
        self.assertFalse(wo.cost_lines.filter(description="Boya").exists())
        self.client.post(
            preview_url,
            {"sheet": "0", "kind-0": "material", "use": ["0:0", "0:1"]},
        )
        names = set(wo.cost_lines.filter(kind=CostLineKind.MATERIAL).values_list("description", flat=True))
        self.assertEqual(names, {"Köhnə boya", "NASOS — Nasos", "Boya", "Fırça"})
        brush = wo.cost_lines.get(description="Fırça")
        self.assertEqual(brush.qty, Decimal("1.000"))
        self.assertEqual(brush.unit, "ədəd")
        self.assertEqual(brush.unit_price, Decimal("0.00"))
        locked.refresh_from_db()
        self.assertEqual(locked.amount, Decimal("8.00"))

        labor = self._xlsx([["Meyar", "Miqdar", "Vahid", "Qiymət"], ["Usta", 3, "", 20]])
        self.client.post(detail, {"action": "import_labor", "labor_file": labor})
        response = self.client.post(
            preview_url,
            {"sheet": "0", "kind-0": "labor", "use": ["0:0"]},
            follow=True,
        )
        self.assertContains(response, "1 sətir əlavə olundu.")
        job = wo.cost_lines.get(kind=CostLineKind.LABOR)
        self.assertEqual(job.unit, "saat")
        self.assertEqual(job.amount, Decimal("60.00"))

        before = wo.cost_lines.count()
        bad = self._xlsx([["Ad", "Miqdar", "Vahid", "Qiymət"], ["Səhv", "abc", "ədəd", 1]])
        response = self.client.post(detail, {"action": "import_material", "material_file": bad}, follow=True)
        self.assertContains(response, "Sətir 2: miqdar rəqəm olmalıdır.")
        self.assertEqual(wo.cost_lines.count(), before)

        template = self.client.get(reverse("erp:work_order_sheet_template", args=[wo.code]) + "?kind=labor")
        self.assertEqual(template.status_code, 200)
        book = load_workbook(BytesIO(template.content))
        self.assertEqual(book.active["A1"].value, "Ad")
        self.assertEqual(book.active["A2"].value, "Usta işi")

    def test_smeta_preview_keeps_one_sheet_and_skips_totals(self):
        self.client.force_login(self.desk)
        self.client.post(reverse("erp:ticket_detail", args=[self.ticket.code]), {"action": "create_wo"})
        wo = WorkOrder.objects.get(ticket=self.ticket)
        kept = WorkOrderCostLine.objects.create(
            work_order=wo,
            kind=CostLineKind.MATERIAL,
            source=CostLineSource.MANUAL,
            description="Köhnə boya",
            qty=Decimal("1"),
            unit="ədəd",
            unit_price=Decimal("3"),
            amount=Decimal("3"),
        )
        locked = WorkOrderCostLine.objects.create(
            work_order=wo,
            kind=CostLineKind.MATERIAL,
            source=CostLineSource.WAREHOUSE,
            description="NASOS — Nasos",
            qty=Decimal("2"),
            unit="ədəd",
            unit_price=Decimal("4"),
            amount=Decimal("8"),
        )
        book = Workbook()
        first = book.active
        first.title = "1"
        first.append(["Elektrik"])
        first.append(["№", "Kontragent", "Malın adı", "Ölcü vahidi", "Miqdar", "Qiymət", "Cəm", "Qeyd"])
        first.append([1, "-", "Kabel", "m", 10, 8.1999999999999993, 82])
        first.append([2, "-", "Boru", "m", 2, None, 24])
        first.append([106])
        first.append(["İşçilik"])
        first.append(["№", "Kontragent", "İşin adı", "Ölcü vahidi", "Miqdar", "Qiymət", "Ödəniş", "Qeyd"])
        first.append([1, "-", "Montaj", "kv", 2, 10, 20])
        first.append([20])
        first.append(["Qeyd: son qiymət fərqlənə bilər", None, None, None, None, None, "Cəm", 126])
        other = book.create_sheet("Maya")
        other.append(["Elektrik"])
        other.append(["№", "Kontragent", "Malın adı", "Ölcü vahidi", "Miqdar", "Qiymət", "Cəm"])
        other.append([1, "-", "Başqa", "ədəd", 1, 5, 5])
        buffer = BytesIO()
        book.save(buffer)
        upload = SimpleUploadedFile(
            "smeta.xlsx",
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        detail = reverse("erp:work_order_detail", args=[wo.code])
        preview_url = reverse("erp:work_order_import", args=[wo.code])
        self.client.post(detail, {"action": "import_material", "material_file": upload})
        page = self.client.get(preview_url)
        self.assertContains(page, "Kabel")
        self.assertContains(page, "Montaj")
        self.assertNotContains(page, "Başqa")
        other_sheet = self.client.get(preview_url + "?sheet=1")
        self.assertContains(other_sheet, "Başqa")
        self.client.post(
            preview_url,
            {
                "sheet": "0",
                "kind-0": "material",
                "kind-1": "labor",
                "use": ["0:0", "0:1", "1:0"],
            },
        )
        names = set(wo.cost_lines.values_list("description", flat=True))
        self.assertIn("Elektrik — Kabel", names)
        self.assertIn("Elektrik — Boru", names)
        self.assertIn("İşçilik — Montaj", names)
        self.assertNotIn("Başqa", names)
        self.assertNotIn("Qeyd: son qiymət fərqlənə bilər", names)
        cable = wo.cost_lines.get(description="Elektrik — Kabel")
        self.assertEqual(cable.kind, CostLineKind.MATERIAL)
        self.assertEqual(cable.unit_price, Decimal("8.20"))
        pipe = wo.cost_lines.get(description="Elektrik — Boru")
        self.assertEqual(pipe.unit_price, Decimal("12.00"))
        job = wo.cost_lines.get(description="İşçilik — Montaj")
        self.assertEqual(job.kind, CostLineKind.LABOR)
        kept.refresh_from_db()
        locked.refresh_from_db()
        self.assertEqual(kept.amount, Decimal("3"))
        self.assertEqual(locked.amount, Decimal("8"))
