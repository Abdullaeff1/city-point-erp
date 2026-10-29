import io
from datetime import date

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse

from apps.accounts.models import Role
from apps.property.models import Building, BuildingAsset, BuildingAssetKind, Floor, add_one_month

User = get_user_model()


class BuildingAssetPageTests(TestCase):
    def setUp(self):
        self.fm = User.objects.create_user(
            username="fm-assets",
            email="fm-assets@test.az",
            password="Pass12345!",
            role=Role.PROPERTY_FM,
            is_staff=True,
        )
        self.desk = User.objects.create_user(
            username="desk-assets",
            email="desk-assets@test.az",
            password="Pass12345!",
            role=Role.SERVICE_DESK,
            is_staff=True,
        )
        building = Building.objects.create(code="CP-BA", name="City Point")
        self.floor = Floor.objects.create(building=building, code="F1", name="1-ci mərtəbə")
        self.client = Client()

    def test_fm_adds_equipment_with_monthly_inspection(self):
        self.client.force_login(self.fm)
        response = self.client.post(
            reverse("erp:building_equipment"),
            {
                "name": "Lift 1",
                "category": "lift",
                "floor": self.floor.pk,
                "location": "Hol",
                "quantity": "2",
                "last_inspected_on": "2026-01-15",
            },
        )
        self.assertEqual(response.status_code, 302)
        asset = BuildingAsset.objects.get(name="Lift 1")
        self.assertEqual(asset.kind, BuildingAssetKind.EQUIPMENT)
        self.assertEqual(asset.category, "lift")
        self.assertEqual(asset.floor_id, self.floor.pk)
        self.assertEqual(asset.quantity, 2)
        self.assertEqual(asset.next_inspection_on, date(2026, 2, 15))
        page = self.client.get(reverse("erp:building_equipment"))
        self.assertContains(page, "Lift 1")
        self.assertContains(page, "Avadanlıqlar")
        inventory = self.client.get(reverse("erp:building_inventory"))
        self.assertNotContains(inventory, "Lift 1")
        updated = self.client.post(
            reverse("erp:building_equipment"),
            {
                "action": "update",
                "asset_id": asset.pk,
                "name": "Lift A",
                "location": "Blok A",
                "last_inspected_on": "2026-03-10",
                "next_inspection_on": "2026-04-10",
            },
        )
        self.assertEqual(updated.status_code, 302)
        asset.refresh_from_db()
        self.assertEqual(asset.name, "Lift A")
        self.assertEqual(asset.location, "Blok A")
        self.assertEqual(asset.next_inspection_on, date(2026, 4, 10))

    def test_inventory_page_separate(self):
        self.client.force_login(self.fm)
        response = self.client.post(
            reverse("erp:building_inventory"),
            {
                "name": "Ofis masası",
                "category": "furniture",
                "floor": self.floor.pk,
                "quantity": "4",
            },
        )
        self.assertEqual(response.status_code, 302)
        asset = BuildingAsset.objects.get(name="Ofis masası")
        self.assertEqual(asset.kind, BuildingAssetKind.INVENTORY)
        page = self.client.get(reverse("erp:building_inventory"))
        self.assertContains(page, "Ofis masası")
        self.assertContains(page, "İnventarlar")
        self.assertNotContains(self.client.get(reverse("erp:building_equipment")), "Ofis masası")

    def test_service_desk_opens_both_pages(self):
        self.client.force_login(self.desk)
        equipment = self.client.get(reverse("erp:building_equipment"))
        self.assertEqual(equipment.status_code, 200)
        self.assertContains(equipment, "Avadanlıqlar")
        self.assertContains(equipment, 'href="/erp/building-equipment/"')
        self.assertContains(equipment, 'href="/erp/building-inventory/"')
        inventory = self.client.get(reverse("erp:building_inventory"))
        self.assertEqual(inventory.status_code, 200)
        self.assertContains(inventory, "İnventarlar")
        legacy = self.client.get(reverse("erp:building_assets"))
        self.assertEqual(legacy.status_code, 302)
        self.assertEqual(legacy.url, reverse("erp:building_equipment"))

    def test_add_one_month_clamps_short_month(self):
        self.assertEqual(add_one_month(date(2026, 1, 31)), date(2026, 2, 28))
        self.assertEqual(add_one_month(date(2024, 1, 31)), date(2024, 2, 29))

    def _xlsx(self):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Ad", "Kod", "Kateqoriya", "Mərtəbə", "Yer", "Say", "Son baxış", "Qeyd"])
        sheet.append(["Lift 1", "L-1", "Lift", "1", "Hol", 2, date(2026, 1, 15), ""])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return SimpleUploadedFile(
            "avadanliq.xlsx",
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_xlsx_import_sets_next_inspection(self):
        self.client.force_login(self.fm)
        preview = self.client.post(
            reverse("erp:building_equipment"),
            {"action": "import", "asset_file": self._xlsx()},
            follow=True,
        )
        self.assertContains(preview, "Lift 1")
        self.assertContains(preview, "Siyahıya yaz")
        saved = self.client.post(
            reverse("erp:building_equipment"), {"action": "import_save"}, follow=True
        )
        self.assertEqual(saved.status_code, 200)
        asset = BuildingAsset.objects.get(name="Lift 1", code="L-1")
        self.assertEqual(asset.kind, BuildingAssetKind.EQUIPMENT)
        self.assertEqual(asset.category, "lift")
        self.assertEqual(asset.floor_id, self.floor.pk)
        self.assertEqual(asset.location, "Hol")
        self.assertEqual(asset.quantity, 2)
        self.assertEqual(asset.last_inspected_on, date(2026, 1, 15))
        self.assertEqual(asset.next_inspection_on, date(2026, 2, 15))
        again = self.client.post(
            reverse("erp:building_equipment"),
            {"action": "import", "asset_file": self._xlsx()},
            follow=True,
        )
        self.assertEqual(again.status_code, 200)
        self.client.post(reverse("erp:building_equipment"), {"action": "import_save"})
        self.assertEqual(BuildingAsset.objects.filter(name="Lift 1", code="L-1").count(), 1)

    def test_docx_line_imports_by_name(self):
        from docx import Document

        self.client.force_login(self.fm)
        document = Document()
        document.add_paragraph("Nasos otağı")
        buffer = io.BytesIO()
        document.save(buffer)
        uploaded = SimpleUploadedFile(
            "avadanliq.docx",
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.client.post(
            reverse("erp:building_equipment"),
            {"action": "import", "asset_file": uploaded},
        )
        self.client.post(reverse("erp:building_equipment"), {"action": "import_save"})
        asset = BuildingAsset.objects.get(name="Nasos otağı")
        self.assertEqual(asset.kind, BuildingAssetKind.EQUIPMENT)
        self.assertEqual(asset.category, "other")

    def test_service_desk_imports_xlsx(self):
        self.client.force_login(self.desk)
        response = self.client.post(
            reverse("erp:building_equipment"),
            {"action": "import", "asset_file": self._xlsx()},
            follow=True,
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Lift 1")
        self.client.post(reverse("erp:building_equipment"), {"action": "import_save"})
        self.assertTrue(BuildingAsset.objects.filter(name="Lift 1", code="L-1").exists())
